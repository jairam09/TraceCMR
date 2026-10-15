"""Smoke test of the napari panel: open a scan, compute, check the numbers.

Skipped unless napari and a display (real or Xvfb) are available.
"""

import os

import numpy as np
import pytest

napari = pytest.importorskip("napari")
if not os.environ.get("DISPLAY") and os.environ.get("QT_QPA_PLATFORM") != "offscreen":
    pytest.skip("needs a display (run under xvfb-run)", allow_module_level=True)

from tracecmr.app._widget import TraceCMRWidget  # noqa: E402
from tracecmr.app.session import Session  # noqa: E402
from tracecmr.synthetic import tapered_phantom  # noqa: E402
from qtpy.QtCore import Qt  # noqa: E402
from qtpy.QtWidgets import QApplication  # noqa: E402


def test_panel_computes_and_reorders():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        assert panel.list.topLevelItemCount() == ed.image.shape[2]
        panel.compute()
        assert float(panel.res_labels["ef_pct"].text()) > 0
        panel._flip()
        assert "reversed" in panel.lbl_order.text()
        # an edit in the napari layer reaches the results
        ef_before = panel.res_labels["edv_ul"].text()
        data = viewer.layers["contours"].data
        data[0][data[0] == 1] = 2  # turn ED blood pool into myocardium on every slice
        viewer.layers["contours"].data = data
        panel.compute()
        assert "need a drawn blood pool" in panel.lbl_notes.text() or panel.res_labels["edv_ul"].text() != ef_before
    finally:
        viewer.close()


def test_dragging_a_control_point_changes_volume():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        panel.compute()
        edv0 = float(panel.res_labels["edv_ul"].text())
        pts = viewer.layers["control points"]
        data = np.asarray(pts.data).copy()
        idx = [i for i, k in enumerate(panel.editor._keys) if k[0] == 0 and k[2] == "endo"]
        z = panel.editor._keys[idx[0]][1]
        same = [i for i in idx if panel.editor._keys[i][1] == z]
        c = data[same, 2:].mean(0)
        for i in same:  # enlarge one slice's endocardium by 30%
            data[i, 2:] = c + 1.3 * (data[i, 2:] - c)
        pts.data = data
        assert (0, z, "endo") in panel.editor.edited
        panel.compute()
        assert float(panel.res_labels["edv_ul"].text()) > edv0
        panel.rb_pixels.setChecked(True)
        assert viewer.layers["contours"].visible and not viewer.layers["control points"].visible
    finally:
        viewer.close()


def test_set_ed_es_from_current_frame():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    img = np.stack([ed.image, ed.image, es.image, es.image], -1)
    lab = np.stack([ed.label, np.zeros_like(ed.label), es.label, np.zeros_like(es.label)], -1)
    s = Session(image=img, label=lab, spacing=ed.spacing, subject_id="phantom", ed_frame=0, es_frame=0)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        panel._show_frame(2)
        panel.btn_set_es.click()
        assert panel.session.es_frame == 2 and "ES" in panel.lbl_frame.text()
        panel._show_frame(0)
        panel.btn_set_ed.click()
        assert panel.session.ed_frame == 0
        panel.compute()
        assert float(panel.res_labels["ef_pct"].text()) > 0
    finally:
        viewer.close()


def test_autocontour_button_with_a_stand_in_model(monkeypatch):
    """The real models are optional installs; a stand-in returns phantom contours on every frame."""
    from tracecmr import autocontour as ac

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    img = np.stack([ed.image, ed.image, es.image, ed.image], -1)
    truth = np.stack([ed.label, ed.label, es.label, ed.label], -1)

    def fake_run(image, spacing, model="ucl", frames=None, progress=None):
        out = np.zeros(image.shape, np.uint8)
        for t in frames:
            out[..., t] = truth[..., t]
        out[:, :, 3, 0] = 0  # a missed slice on frame 0, so sanity checks have something to flag
        return out

    monkeypatch.setattr(ac, "available", lambda: {"ucl": "ready", "calico": "not installed"})
    monkeypatch.setattr(ac, "run", fake_run)
    s = Session(image=img, label=np.zeros(img.shape, np.uint8), spacing=ed.spacing, subject_id="raw")
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        panel.chk_all_frames.setChecked(True)
        panel.autocontour(sync=True)
        assert "Contoured 4 frame(s)" in panel.lbl_auto.text()
        assert panel.session.ed_frame in (1, 3) and panel.session.es_frame == 2  # from the volume curve (frame 0 lost a slice)
        assert len(panel.vfig.axes) == 1  # volume curve drawn
        panel.compute()
        assert float(panel.res_labels["ef_pct"].text()) > 0
    finally:
        viewer.close()


def test_smooth_button_and_undo():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        ed_ = panel.editor
        z = int(viewer.dims.current_step[1])
        before = {k: v.copy() for k, v in ed_.curves.items()}
        n_points = len(ed_.points.data)
        panel._smooth_points()  # Endo is selected: only the endocardium is smoothed
        here = {(0, z, "endo")}
        assert ed_.edited == here and not ed_.hidden
        assert len(ed_.points.data) == n_points  # its points stay, ready to adjust
        assert all(np.allclose(before[k], ed_.curves[k]) for k in before if k not in here)
        # Points off: a clean view (no dots at all); picking a contour (Q) brings them back
        panel.btn_points.setChecked(False)
        assert len(ed_.points.data) == 0
        panel.cycle_structure()
        assert panel.btn_points.isChecked() and len(ed_.points.data) > 0
        panel.btn_points.click()
        assert len(ed_.points.data) == 0
        panel.btn_points.click()
        assert len(ed_.points.data) > 0
        # editing another slice keeps the smoothed slice's contours
        other = next(k for k in ed_.curves if k[:2] != (0, z))
        ed_.move_point(other, 0, ed_.curves[other][0] + 1.0)
        assert all(k in ed_.curves for k in here)
        assert ed_.undo()
    finally:
        viewer.close()


def test_smoothed_contour_stays_smoothed_after_phase_change_elsewhere():
    """Smooth on one slice, set ES on another slice, come back: the contours stay smoothed."""
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image, es.image], -1),
                label=np.stack([ed.label, es.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        z = int(viewer.dims.current_step[1])
        for target in (0, 1):  # smooth the endo, then the epi
            panel.cmb_add.setCurrentIndex(target)
            panel._smooth_points()
        smoothed = {k: e.curves[k].copy() for k in ((0, z, "endo"), (0, z, "epi"))}
        panel.chk_per_slice.setChecked(True)
        panel.step_slice(1)
        panel.set_phase_frame("ES", 2)  # re-shows the curves
        panel.step_slice(-1)
        assert all(np.allclose(e.curves[k], v) for k, v in smoothed.items())
        assert len(e.points.data) == len(smoothed[(0, z, "epi")])  # epi selected: its (smoothed) points show
        # contours replaced on this slice (e.g. auto-contour) come back with dots
        lab = panel.session.label.copy()
        lab[:, :, z, 0] = 0
        lab[:, :, z, 0][ed.label[:, :, z] > 0] = ed.label[:, :, z][ed.label[:, :, z] > 0]
        lab[:, :, z, 0] = np.roll(lab[:, :, z, 0], 1, axis=0)
        panel.session.label = lab
        panel.viewer.layers["contours"].data = lab.transpose(3, 2, 1, 0).copy()
        panel._mode_changed()
        assert not any(k[:2] == (0, z) for k in e.hidden)
    finally:
        viewer.close()


def test_zoom_buttons_and_fit():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    s = Session(image=np.stack([ed.image, ed.image], -1), label=np.stack([ed.label, ed.label], -1),
                spacing=ed.spacing, subject_id="phantom", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        z0 = viewer.scene.camera.zoom
        panel.btn_zoom_in.click()
        assert viewer.scene.camera.zoom == pytest.approx(z0 * 1.25)
        panel.btn_zoom_out.click()
        panel.btn_zoom_out.click()
        assert viewer.scene.camera.zoom < z0
        step = viewer.dims.current_step
        panel.btn_zoom_fit.click()  # back to the heart view, same slice
        assert viewer.scene.camera.zoom == pytest.approx(z0, rel=0.01)
        assert viewer.dims.current_step == step
    finally:
        viewer.close()


def test_model_folder_choice(tmp_path, monkeypatch):
    monkeypatch.delenv("TRACECMR_UCL_DIR", raising=False)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        status = panel.set_model_dir("ucl", str(tmp_path), remember=False)  # empty folder: not a model
        assert status != "ready" and "not found" in panel.lbl_auto.text()
        assert os.environ["TRACECMR_UCL_DIR"] == str(tmp_path)
        assert "(not installed)" in panel.cmb_model.currentText()
    finally:
        viewer.close()



class _Ev:
    """Minimal stand-in for a napari mouse event."""

    def __init__(self, world, button=1, screen=(0.0, 0.0)):
        self.position, self.button, self.pos, self.type = tuple(world), button, screen, "mouse_press"


def _drive(editor, viewer, ev, moves=()):
    gen = editor._on_mouse(viewer, ev)
    try:
        next(gen)
        for world, screen in moves:
            ev.type, ev.position, ev.pos = "mouse_move", tuple(world), screen
            next(gen)
        ev.type = "mouse_release"
        next(gen)
    except StopIteration:
        pass


def _world(editor, t, z, p):
    sc = editor.scale
    return (t, z * sc[1], p[1] * sc[2], p[0] * sc[3])


def test_segment_style_mouse_editing_and_undo():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        key = (t, z, "endo")
        n0 = len(e.curves[key])
        p0 = e.curves[key][0].copy()
        # drag the first dot outward
        c = e.curves[key].mean(0)
        target = c + 1.3 * (p0 - c)
        _drive(e, viewer, _Ev(_world(e, t, z, p0)), [(_world(e, t, z, target), (5, 5))])
        assert np.allclose(e.curves[key][0], target) and key in e.edited
        # click on the line between dots 2 and 3: adds a dot
        mid = spline_mid = None
        from tracecmr.contours import spline_through
        dense = spline_through(e.curves[key], 400)
        per = max(2, 400 // len(e.curves[key]))
        spline_mid = dense[2 * per + per // 2]
        _drive(e, viewer, _Ev(_world(e, t, z, spline_mid)))
        assert len(e.curves[key]) == n0 + 1
        # right-click a dot deletes it
        _drive(e, viewer, _Ev(_world(e, t, z, e.curves[key][5]), button=2))
        assert len(e.curves[key]) == n0
        # three undos bring the original contour back
        for _ in range(3):
            panel.undo()
        assert len(e.curves[key]) == n0 and np.allclose(e.curves[key][0], p0)
        # trash can clears the selected contour (Endo), then undo
        panel.clear_slice()
        assert key not in e.curves and (t, z, "epi") in e.curves
        panel.undo()
        assert key in e.curves
        # draw a new endocardium on a cleared slice: 4 clicks, right-click closes
        panel.clear_slice()
        ring = [c + 4 * np.array([np.cos(a), np.sin(a)]) for a in np.linspace(0, 2 * np.pi, 5)[:-1]]
        for q in ring:
            _drive(e, viewer, _Ev(_world(e, t, z, q)))
        assert e.pending is not None and len(e.pending[3]) == 4
        _drive(e, viewer, _Ev(_world(e, t, z, c + 30), button=2))
        assert key in e.curves and len(e.curves[key]) == 4 and e.pending is None
        panel.compute()  # the drawn contour feeds the numbers
        assert float(panel.res_labels["edv_ul"].text()) > 0
    finally:
        viewer.close()


def test_right_drag_pans_without_editing():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    s = Session(image=np.stack([ed.image, ed.image], -1), label=np.stack([ed.label, ed.label], -1),
                spacing=ed.spacing, subject_id="phantom", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        before = {k: v.copy() for k, v in e.curves.items()}
        centre = viewer.scene.camera.center
        p = e.curves[(t, z, "endo")][0]
        _drive(e, viewer, _Ev(_world(e, t, z, p), button=2, screen=(100, 100)),
               [(_world(e, t, z, p), (140, 100)), (_world(e, t, z, p), (180, 120))])
        assert viewer.scene.camera.center != centre
        assert all(np.array_equal(before[k], e.curves[k]) for k in before)
        assert viewer.scene.camera.mouse_pan is False  # left-drag edits, not pans
    finally:
        viewer.close()


def test_left_drag_never_pans_unless_pan_is_on():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    s = Session(image=np.stack([ed.image, ed.image], -1), label=np.stack([ed.label, ed.label], -1),
                spacing=ed.spacing, subject_id="phantom", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        centre = viewer.scene.camera.center
        corner = np.array([1.0, 1.0])  # far from the heart
        _drive(e, viewer, _Ev(_world(e, t, z, corner), screen=(10, 10)), [(_world(e, t, z, corner), (60, 40))])
        assert viewer.scene.camera.center == centre and e.pending is None  # editing never pans
        panel.btn_pan.setChecked(True)  # Pan on: left-drag moves the image
        _drive(e, viewer, _Ev(_world(e, t, z, corner), screen=(10, 10)), [(_world(e, t, z, corner), (60, 40))])
        assert viewer.scene.camera.center != centre
        panel.btn_pan.setChecked(False)
        centre = viewer.scene.camera.center
        # a drag that starts a little off a dot (beyond the click radius, inside the grab radius) moves that dot
        key = (t, z, "endo")
        p0 = e.curves[key][0].copy()
        c = e.curves[key].mean(0)
        off = p0 + (p0 - c) / np.linalg.norm(p0 - c) * e._tol_px() * 1.6
        _drive(e, viewer, _Ev(_world(e, t, z, off), screen=(10, 10)), [(_world(e, t, z, off + 2), (30, 30))])
        assert not np.allclose(e.curves[key][0], p0) and viewer.scene.camera.center == centre  # no pan

        # drawing: a misplaced point can be dragged before the contour is closed
        panel.clear_slice()  # endo gone on this slice
        ring = [c + 4 * np.array([np.cos(a), np.sin(a)]) for a in np.linspace(0, 2 * np.pi, 5)[:-1]]
        for q in ring:
            _drive(e, viewer, _Ev(_world(e, t, z, q)))
        first = np.array(e.pending[3][0])
        _drive(e, viewer, _Ev(_world(e, t, z, first)), [(_world(e, t, z, first + 1.5), (20, 20))])
        assert np.allclose(e.pending[3][0], first + 1.5) and len(e.pending[3]) == 4
    finally:
        viewer.close()


def test_series_panel_lists_study_and_switches(tmp_path):
    from test_study import _write

    _write(tmp_path / "a", "cine_sax", 3, zs=range(4))
    _write(tmp_path / "b", "cine_4ch", 4, zs=[0], orient=[1, 0, 0, 0, 0, 1])
    for k in range(3):  # one cine per slice, as in Segment workflows
        _write(tmp_path / f"s{k}", f"SAX_S-{k}", 10 + k, zs=[[2, 0, 1][k]])
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.open_study(tmp_path)
        # stack row, the multi-slice SAX and the 4CH (the stack's slices are in the slice list)
        assert panel.series_panel.list.count() == 3
        assert panel.list.topLevelItemCount() == 3
        s = panel.session
        assert s.image.shape[2] == 3 and s.separate_slices and panel.chk_per_slice.isChecked()
        assert s.slice_names == ["11 SAX_S-1", "12 SAX_S-2", "10 SAX_S-0"]  # sorted by slice position
        # clicking a slice row shows that slice of the stack
        part = next(x for x in panel.study if x.stack is not None and x.slice_index == 2)
        panel.open_series(panel.study.index(part))
        assert panel.current_slice() == 2
        # per-slice ED/ES: set ES on this slice only
        panel._show_frame(3)
        panel.btn_set_es.click()
        assert s.t_for(2, "ES") == 3 and s.t_for(0, "ES") == 0
        panel.timebar.grab()  # paints without error
        # an edit survives switching away and back
        viewer.layers["contours"].data[0, 1, 2, 2] = 1
        stack_i = panel.current_series
        lax = next(i for i, x in enumerate(panel.study) if x.kind == "Long axis")
        panel.open_series(lax)
        assert panel.session.image.shape[2] == 1
        panel.open_series(stack_i)
        assert panel.session.label[2, 2, 1, 0] == 1 and panel.session.t_for(2, "ES") == 3
    finally:
        viewer.close()



def test_autocontour_only_the_frame_being_viewed(monkeypatch):
    """Default: ED = ES = frame 0 until set; Auto-contour does the viewed frame only (one model pass)."""
    from tracecmr import autocontour as ac

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    img = np.stack([ed.image, ed.image, es.image, ed.image], -1)
    truth = np.stack([ed.label, ed.label, es.label, ed.label], -1)
    calls = []

    def fake_pairs(image, spacing, pairs, model="ucl", unsure_out=None):
        calls.append(list(pairs))
        return np.stack([truth[:, :, z, t] for z, t in pairs], axis=2)

    monkeypatch.setattr(ac, "available", lambda: {"ucl": "ready", "calico": "not installed"})
    monkeypatch.setattr(ac, "run_pairs", fake_pairs)
    s = Session(image=img, label=np.zeros(img.shape, np.uint8), spacing=ed.spacing, subject_id="raw")
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(s)
        assert s.ed_frame == 0 and s.es_frame == 0
        panel._show_frame(2)
        panel.autocontour(sync=True)
        assert calls == [[(z, 2) for z in range(img.shape[2])]]  # every slice, only frame 2
        assert s.label[..., 2].any() and not s.label[..., [0, 1, 3]].any()
        assert panel.current_frame() == 2  # stays on the frame just contoured
        # per-slice mode at a slice's ES frame: each slice's own ES
        panel.chk_per_slice.setChecked(True)
        s.set_phase_frame("ES", 2)
        s.set_phase_frame("ES", 3, z=1)
        assert panel.autocontour_pairs()[1] == (1, 3) and panel.autocontour_pairs()[0] == (0, 2)
    finally:
        viewer.close()


def test_hand_drawn_rv_and_rv_volumes():
    from tracecmr import labels as L

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    lab = np.stack([ed.label, es.label], -1)
    lab[2:10, 2:12, 1:7, 0] = L.RV_BLOODPOOL  # an RV beside the LV, bigger at ED
    lab[2:8, 2:12, 1:7, 1] = L.RV_BLOODPOOL
    s = Session(image=np.stack([ed.image, es.image], -1), label=lab, spacing=ed.spacing, subject_id="rv",
                papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        z = int(viewer.dims.current_step[1])
        assert (0, z, "rv") in panel.editor.curves  # yellow RV curve is editable
        panel.btn_rv.click()
        assert set(k[2] for k in panel.editor._keys) == {"rv"}
        panel.compute()
        assert panel.box_rv.isVisibleTo(panel) and float(panel.rv_labels["rvef_pct"].text()) > 0
        assert not hasattr(panel, "btn_auto_rv")  # no RV model in the app
    finally:
        viewer.close()

def test_contours_show_and_edit_on_any_frame():
    """Contours on a frame that is neither ED nor ES (Auto LV on every frame, or drawn there) show up as editable
    curves when that frame is viewed, edits reach the labels, and undo never clears them."""
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    mid = tapered_phantom(n_slices=8, base_blood_mm=2.0, apex_blood_mm=0.8, wall_mm=1.1)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, mid.image, es.image], -1),
                label=np.stack([ed.label, mid.label, es.label], -1), spacing=ed.spacing, subject_id="phantom",
                papillary="blood_pool", ed_frame=0, es_frame=2)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        z = int(viewer.dims.current_step[1])
        assert not any(k[0] == 1 for k in e.curves)  # frame 1 isn't built until it is viewed
        e.push_undo()  # an earlier edit on ED, to undo after frame 1 has loaded
        e.move_point((0, z, "endo"), 0, e.curves[(0, z, "endo")][0] + 0.5)
        viewer.dims.set_current_step(0, 1)
        assert (1, z, "endo") in e.curves and (1, z, "epi") in e.curves
        assert any(k[0] == 1 for k in e._keys)  # its dots are on screen
        before = s.label[:, :, z, 1].copy()
        e.undo()  # undoes the ED edit; frame 1's contours stay
        assert (1, z, "endo") in e.curves
        panel._sync_labels()
        assert np.array_equal(s.label[:, :, z, 1], before)
        e.push_undo()
        pts = e.curves[(1, z, "endo")]
        e.move_point((1, z, "endo"), 0, pts[0] + (pts[0] - pts.mean(axis=0)) * 0.5)  # pull one dot outwards
        panel._sync_labels()
        assert not np.array_equal(s.label[:, :, z, 1], before)
    finally:
        viewer.close()


def test_slice_list_keeps_room_on_a_short_screen():
    """On a short screen the mice and series lists shrink (they scroll) so the slice list still shows up to
    6 slices, instead of being pushed below the bottom of the panel."""
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        sp = panel.series_panel
        for i in range(14):
            sp.mice.addItem(f"Mouse{i}.h5")
        for i in range(9):
            sp.list.addItem(f"series {i}")
        sp.avail_height = lambda: 2000  # a tall screen: everything at its usual size
        sp._fit_lists()
        assert sp.mice.height() == 5 * sp.MICE_ROW + 6 and sp.list.height() == 3 * sp.SERIES_ROW + 6
        assert panel.list.minimumHeight() == 6 * sp.SLICE_ROW + 30
        sp.avail_height = lambda: 420  # a short one
        sp._fit_lists()
        sp.layout().activate()
        assert sp.mice.height() < 5 * sp.MICE_ROW + 6 and sp.list.height() <= 2 * sp.SERIES_ROW + 6
        assert panel.list.minimumHeight() == 6 * sp.SLICE_ROW + 30  # the slices keep their room
        assert sp.layout().minimumSize().height() <= 420 or (sp.mice.height(), sp.list.height()) == (30, 60)
    finally:
        viewer.close()


def test_open_scan_starts_in_the_last_folder(qtbot, monkeypatch, tmp_path):
    """Open scan's folder picker opens where the last scan came from (its parent folder), not in Documents."""
    from tracecmr.app import _widget as W

    store = {}
    monkeypatch.setattr(W, "_settings", lambda: type("S", (), {"value": lambda self, k, d=None: store.get(k, d),
                                                                 "setValue": lambda self, k, v: store.update({k: v})})())
    scans = tmp_path / "mri"
    (scans / "study_44").mkdir(parents=True)
    starts, opened = [], []

    def pick(parent, caption, start=""):
        starts.append(start)
        return str(scans / "study_44")

    monkeypatch.setattr(W.QFileDialog, "getExistingDirectory", staticmethod(pick))
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        monkeypatch.setattr(panel, "open_path", opened.append)
        panel._open_dicom()
        panel._open_dicom()
        assert starts[0] == str(W.Path.home()) and starts[1] == str(scans)
        assert opened == [str(scans / "study_44")] * 2
    finally:
        viewer.close()


def test_contour_focus_and_per_slice_results():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        assert set(k[2] for k in e._keys) == {"endo"}  # only the selected contour's dots are shown
        panel.btn_epi.click()
        assert set(k[2] for k in e._keys) == {"epi"} and e.add_to == "epi"
        # with Epi selected, a click on an endocardium dot grabs nothing
        p = e.curves[(t, z, "endo")][0]
        assert e.hit(t, z, p, 0.5) is None or e.hit(t, z, p, 0.5)[1][2] == "epi"
        panel._smooth_points()
        assert e.edited == {(t, z, "epi")}
        # per-slice table: one row per slice with ED and ES, plus the total
        panel.compute()
        rows = s.per_slice()
        assert panel.box_slices.isVisibleTo(panel) and panel.tbl_slices.rowCount() == len(rows) + 1
        assert sum(r["edv_ul"] for r in rows if r["included"]) == pytest.approx(
            float(panel.res_labels["edv_ul"].text()), rel=0.02)
        # leaving a slice out greys it and changes the total
        item = panel.list.topLevelItem(1)
        zz = item.data(0, Qt.ItemDataRole.UserRole)
        item.setCheckState(1, Qt.CheckState.Unchecked)
        panel._refresh_slice_table()
        assert not next(r for r in s.per_slice() if r["slice"] == zz)["included"]
        total = float(panel.tbl_slices.item(panel.tbl_slices.rowCount() - 1, 1).text())
        assert total == pytest.approx(sum(r["edv_ul"] for r in s.per_slice() if r["slice"] in s.included["ED"]), abs=0.1)
    finally:
        viewer.close()


def test_save_and_reload_contours_and_data_folder(tmp_path, monkeypatch):
    from qtpy.QtCore import QSettings
    from test_study import _write

    monkeypatch.setattr(QSettings, "value", lambda self, key, *a, **k: {
        "folders/contours": str(tmp_path / "saved"), "folders/data": str(tmp_path / "data")}.get(key))
    monkeypatch.setattr(QSettings, "setValue", lambda self, *a: None)
    for k in range(3):
        _write(tmp_path / "data" / "MouseA" / f"s{k}", f"SAX_S-{k}", 10 + k, zs=[k])
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        assert panel.series_panel.mice.count() == 1  # MouseA listed from the data folder
        panel.open_path(tmp_path / "data" / "MouseA")
        s = panel.session
        assert s.subject_id.startswith("MouseA") and s.image.shape[2] == 3
        s.label[2:6, 2:6, 1, 0] = 1  # some contour
        panel.viewer.layers["contours"].data = s.label.transpose(3, 2, 1, 0).copy()
        panel._show_frame(3)
        panel.chk_per_slice.setChecked(True)
        panel.set_phase_frame("ES", 3)
        path = panel.save_contours()
        assert path is not None and path.exists()
        assert "✓" in panel.series_panel.mice.item(0).text()
        # reopen: contours and the per-slice ES come back
        panel.open_path(tmp_path / "data" / "MouseA")
        s2 = panel.session
        assert s2 is not s and s2.label[3, 3, 1, 0] == 1 and s2.t_for(panel.current_slice(), "ES") == 3
        assert "Loaded your saved contours" in panel.lbl_auto.text()
        # no slice scrollbar under the image
        assert not viewer.window._qt_viewer.dims.isVisible()
    finally:
        viewer.close()


def test_results_kept_apart_from_raw_scans(tmp_path, monkeypatch):
    """Contours, tables and images go to <results>/<mouse>/, never into the raw scans folder."""
    from qtpy.QtCore import QSettings
    from test_study import _write

    store = {"folders/contours": str(tmp_path / "results"), "folders/data": str(tmp_path / "data")}
    monkeypatch.setattr(QSettings, "value", lambda self, key, *a, **k: store.get(key))
    monkeypatch.setattr(QSettings, "setValue", lambda self, key, v: store.__setitem__(key, v))
    for k in range(3):
        _write(tmp_path / "data" / "MouseA" / f"s{k}", f"SAX_S-{k}", 10 + k, zs=[k])
    raw_before = {p: p.stat().st_mtime for p in (tmp_path / "data").rglob("*")}
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.open_path(tmp_path / "data" / "MouseA")
        s = panel.session
        s.label[4:12, 4:12, :, 0] = 2
        s.label[6:10, 6:10, :, 0] = 1
        panel.viewer.layers["contours"].data = s.label.transpose(3, 2, 1, 0).copy()
        path = panel.save_contours()
        folder = panel.results_dir()
        assert path.parent == folder and folder.parent == tmp_path / "results"
        imgs = sorted(p.name for p in (folder / "images").glob("*.png"))
        assert "montage_ED.png" in imgs and any(n.startswith("slice") for n in imgs)
        # the raw scans are untouched and nothing new appeared among them
        assert {p: p.stat().st_mtime for p in (tmp_path / "data").rglob("*")} == raw_before
        # results inside the raw folder (or vice versa) are refused
        assert not panel.set_results_dir(tmp_path / "data" / "out")
        assert store["folders/contours"] == str(tmp_path / "results")
        # raw view hides the contours without deleting them; H/Raw again brings them back
        panel.btn_raw.setChecked(True)
        assert not viewer.layers["curves"].visible and not viewer.layers["control points"].visible
        assert panel.session.label.any()
        panel.btn_raw.setChecked(False)
        assert viewer.layers["curves"].visible
        # files saved under the old name (VentriMap) still load
        old_name = path.with_name(path.name.replace(".tracecmr.npz", ".ventrimap.npz"))
        path.rename(old_name)
        panel.open_path(tmp_path / "data" / "MouseA")
        assert panel.session.label.any() and "Loaded your saved contours" in panel.lbl_auto.text()
        path = panel._find_saved(panel.session)
        # contours saved by older versions (flat file in the results folder) still load
        legacy = panel._legacy_contours_path(s)
        path.rename(legacy)
        panel.open_path(tmp_path / "data" / "MouseA")
        assert panel.session.label.any() and "Loaded your saved contours" in panel.lbl_auto.text()
    finally:
        viewer.close()


def test_auto_lv_never_overwrites_finished_slices(monkeypatch):
    """Contour slice 2 (with an RV) by hand, go to another slice, Auto LV: slice 2 stays as it was."""
    from tracecmr import autocontour as ac
    from tracecmr import labels as LB

    ed = tapered_phantom(n_slices=6, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    img = np.stack([ed.image, ed.image], -1)
    calls = []

    def fake_pairs(image, spacing, pairs, model="ucl", unsure_out=None):
        calls.append(list(pairs))
        return np.stack([ed.label[:, :, z] for z, t in pairs], axis=2)

    monkeypatch.setattr(ac, "available", lambda: {"ucl": "ready"})
    monkeypatch.setattr(ac, "run_pairs", fake_pairs)
    lab = np.zeros(img.shape, np.uint8)
    mine = ed.label[:, :, 2].copy()
    mine[mine == LB.LV_BLOODPOOL] = LB.LV_MYOCARDIUM  # 'my' slice 2: different from what the model gives
    mine[2:6, 2:6] = LB.RV_BLOODPOOL
    lab[:, :, 2, 1] = mine
    s = Session(image=img, label=lab, spacing=ed.spacing, subject_id="m")
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(s)
        panel._show_frame(1)
        step = list(viewer.dims.current_step); step[1] = 3; viewer.dims.current_step = tuple(step)
        panel.autocontour(sync=True)
        s = panel.session
        assert 2 not in [z for z, _ in calls[-1]]  # slice 2 was not sent to the model
        assert np.array_equal(s.label[:, :, 2, 1], mine)  # and is unchanged, RV included
        assert s.label[:, :, 3, 1].any()
        # every slice done: nothing happens unless the user asks to redo the viewed slice
        n = len(calls)
        panel.autocontour(sync=True, redo_current=False)
        assert len(calls) == n
        step[1] = 2; viewer.dims.current_step = tuple(step)
        panel.autocontour(sync=True, redo_current=True)
        assert calls[-1] == [(2, 1)]
        assert (s.label[:, :, 2, 1] == LB.RV_BLOODPOOL).any()  # hand-drawn RV kept
        # Ctrl+Z reverts the redo
        panel.undo()
        assert np.array_equal(panel.session.label[:, :, 2, 1], mine)
    finally:
        viewer.close()


def test_results_are_live_and_contrast_controls():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        panel._refresh_slice_table()  # what the timer does
        assert panel.res_labels["ef_pct"].text() not in ("–", "")  # no Compute press needed
        layer = viewer.layers["image"]
        base = tuple(layer.contrast_limits)
        panel.sld_contrast.setValue(50)
        c = tuple(layer.contrast_limits)
        assert (c[1] - c[0]) < (base[1] - base[0])  # more contrast = narrower window
        panel.sld_brightness.setValue(50)
        assert tuple(layer.contrast_limits)[0] < c[0]  # brighter = window moves down
        panel.adjust_contrast(10, -10)
        assert panel.sld_contrast.value() == 60 and panel.sld_brightness.value() == 40
        panel._reset_contrast()
        assert np.allclose(layer.contrast_limits, base)
    finally:
        viewer.close()


def test_save_as_asks_once_then_remembers(tmp_path, monkeypatch):
    """First save opens a save window (rename / choose folder); later saves go to the same file; reopening
    the mouse loads from wherever it was saved."""
    from qtpy.QtCore import QSettings
    from qtpy.QtWidgets import QFileDialog
    from test_study import _write

    store = {"folders/contours": str(tmp_path / "results"), "folders/data": str(tmp_path / "data")}
    monkeypatch.setattr(QSettings, "value", lambda self, key, *a, **k: store.get(key))
    monkeypatch.setattr(QSettings, "setValue", lambda self, key, v: store.__setitem__(key, v))
    asked = []
    custom = tmp_path / "elsewhere" / "My mouse A.tracecmr.npz"

    def fake_dialog(*a, **k):
        asked.append(a[2])  # the suggested path
        return str(custom), ""

    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(fake_dialog))
    for k in range(3):
        _write(tmp_path / "data" / "MouseA" / f"s{k}", f"SAX_S-{k}", 10 + k, zs=[k])
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.open_path(tmp_path / "data" / "MouseA")
        s = panel.session
        s.label[4:12, 4:12, :, 0] = 2
        s.label[6:10, 6:10, :, 0] = 1
        panel.viewer.layers["contours"].data = s.label.transpose(3, 2, 1, 0).copy()
        assert panel._save_clicked() == custom and custom.exists() and len(asked) == 1
        assert str(tmp_path / "results") in asked[0]  # suggested: the mouse's results folder
        assert (custom.parent / "My mouse A_images").is_dir()
        panel._save_clicked()  # Ctrl+S again: no window, same file
        assert len(asked) == 1
        panel._save_clicked(save_as=True)  # Save as: asks again
        assert len(asked) == 2
        panel.open_path(tmp_path / "data" / "MouseA")
        assert panel.session.label.any() and custom.name in panel.lbl_auto.text()
    finally:
        viewer.close()


def test_auto_lv_marks_base_and_apex_until_edited(monkeypatch):
    from tracecmr import autocontour as ac

    ed = tapered_phantom(n_slices=6, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    img = np.stack([ed.image, ed.image], -1)
    monkeypatch.setattr(ac, "available", lambda: {"ucl": "ready"})
    monkeypatch.setattr(ac, "run_pairs", lambda image, spacing, pairs, model="ucl", unsure_out=None:
                        np.stack([ed.label[:, :, z] for z, t in pairs], axis=2))
    s = Session(image=img, label=np.zeros(img.shape, np.uint8), spacing=ed.spacing, subject_id="m")
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(s)
        panel.autocontour(sync=True)
        s = panel.session
        zs = [z for z in s.order if s.label[:, :, z, 0].any()]
        base, apex = zs[0], zs[-1]
        assert s.flags[(0, base)].startswith("base slice") and s.flags[(0, apex)].startswith("apex slice")
        assert any(it.text(0).startswith("⚠") for it in panel._items())
        # edit the base slice: its note clears
        e = panel.editor
        key = next(k for k in e.curves if k[:2] == (0, base))
        e.move_point(key, 0, e.curves[key][0] + 0.5)
        panel._sync_labels()
        assert (0, base) not in s.flags and (0, apex) in s.flags
    finally:
        viewer.close()


def test_grow_and_shrink_contour():
    """Grow pushes the selected outline out (never past the epi), Shrink pulls it in, Undo reverts."""
    from tracecmr.contours import polygon_area

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        z = int(viewer.dims.current_step[1])
        endo, epi = (0, z, "endo"), (0, z, "epi")
        a0, epi0 = polygon_area(e.curves[endo]), e.curves[epi].copy()
        panel._smooth_points()
        panel.btn_grow.click()
        a1 = polygon_area(e.curves[endo])
        assert a1 > a0 and endo in e.edited
        assert np.allclose(e.curves[epi], epi0)  # only the selected contour moves
        panel.btn_shrink.click()
        panel.btn_shrink.click()
        assert polygon_area(e.curves[endo]) < a1
        assert e.undo() and e.undo()
        assert np.isclose(polygon_area(e.curves[endo]), a1)
        for _ in range(40):  # never past the epicardium
            panel.grow_contour(0.5)
        from tracecmr.app.curve_editor import _inside
        assert _inside(e.curves[epi], e.curves[endo]).all()
        assert "further" in panel.lbl_auto.text() or "Grew" in panel.lbl_auto.text()
        panel._sync_labels()
        lab = s.label[:, :, z, 0]
        assert (lab == 2).any()  # some myocardium is left
        panel.cmb_add.setCurrentIndex(1)  # Epi: grows outward
        a_epi = polygon_area(e.curves[epi])
        panel.btn_grow.click()
        assert polygon_area(e.curves[epi]) > a_epi
    finally:
        viewer.close()


def test_offset_points_either_winding():
    from tracecmr.app.curve_editor import offset_points
    from tracecmr.contours import polygon_area

    th = np.linspace(0, 2 * np.pi, 20, endpoint=False)
    c = np.column_stack([10 + 5 * np.cos(th), 10 + 5 * np.sin(th)])
    for pts in (c, c[::-1]):
        assert polygon_area(offset_points(pts, 1.0)) > polygon_area(pts)
        assert polygon_area(offset_points(pts, -1.0)) < polygon_area(pts)
        assert np.allclose(np.hypot(*(offset_points(pts, 1.0) - 10).T), 6.0, atol=0.05)


def test_fit_button_pulls_an_oversized_endo_back_to_the_blood():
    from tracecmr.contours import polygon_area

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        z = int(viewer.dims.current_step[1])
        key = (0, z, "endo")
        a0 = polygon_area(e.curves[key])
        for _ in range(6):
            panel.grow_contour(0.5)
        assert polygon_area(e.curves[key]) > a0 * 1.15
        panel.btn_fit.click()
        assert abs(polygon_area(e.curves[key]) - a0) / a0 < 0.12, panel.lbl_auto.text()
        assert "gave" in panel.lbl_auto.text()
        assert e.undo()
        assert polygon_area(e.curves[key]) > a0 * 1.15
    finally:
        viewer.close()


def test_wall_button_draws_thickness_lines_on_the_slice():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        active = viewer.layers.selection.active
        panel.btn_wall.setChecked(True)
        layer = viewer.layers["wall thickness"]
        assert layer.visible and 30 <= len(layer.data) <= 45
        assert viewer.layers.selection.active is active  # editing keeps working
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        lengths = []
        dz, dy, dx = layer.scale[1:]
        for d in layer.data:
            assert int(d[0, 0]) == t and int(d[0, 1]) == z
            lengths.append(np.hypot((d[1, 2] - d[0, 2]) * dy, (d[1, 3] - d[0, 3]) * dx))
        assert np.median(lengths) == pytest.approx(1.0, abs=0.08)  # the phantom wall is 1 mm
        assert "mean" in panel.lbl_auto.text()
        panel.step_slice(1)
        assert all(int(d[0, 1]) != z for d in layer.data)  # follows the slice
        panel.set_wall_mode("thickening")  # phantom: ED wall 1.0 mm, ES wall 1.3 mm -> about +30% all round
        assert 30 <= len(layer.data) <= 45
        import re
        m = re.search(r"mean (-?\d+)%", panel.lbl_auto.text())
        assert m and 22 <= int(m.group(1)) <= 38, panel.lbl_auto.text()
        assert "thinning" not in panel.lbl_auto.text()
        panel.set_wall_mode("thickness")
        assert "mm" in panel.lbl_auto.text() and "Thickening" not in panel.lbl_auto.text()
        panel.btn_wall.setChecked(False)
        assert not layer.visible
    finally:
        viewer.close()


def test_undo_after_live_results_puts_the_labels_back():
    """Fit (or any edit), let the results update (labels synced), then Undo: the labels, results and wall lines
    must follow the undone outline, not keep the edit."""
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        z = int(viewer.dims.current_step[1])
        before = int((s.label[:, :, z, 0] == 1).sum())
        for _ in range(6):
            panel.grow_contour(0.5)
        panel.compute()  # live results: the edit is written into the labels
        grown = int((panel.session.label[:, :, z, 0] == 1).sum())
        assert grown > before * 1.1
        edv_grown = panel.last[0].as_row()["edv_ul"]
        for _ in range(6):
            panel.undo()
        panel.compute()
        after = int((panel.session.label[:, :, z, 0] == 1).sum())
        assert abs(after - before) <= max(4, 0.03 * before), (before, grown, after)
        assert panel.last[0].as_row()["edv_ul"] < edv_grown
    finally:
        viewer.close()


def test_set_septum_by_click_without_an_rv(tmp_path):
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        panel.compute()
        assert not panel.last[1].ed.reference.startswith("septum")  # no RV: image axis
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        panel.start_set_septum()
        assert panel.btn_wall.isChecked() and panel.editor.click_hook is not None
        from scipy import ndimage as ndi

        cx, cy = ndi.center_of_mass(s.label[:, :, z, t] == 1)
        before = {k: v.copy() for k, v in panel.editor.curves.items()}
        _drive(panel.editor, viewer, _Ev(_world(panel.editor, t, z, (cx - 15, cy))))  # a real click, left of the cavity
        assert panel.editor.click_hook is None
        assert all(np.array_equal(before[k], panel.editor.curves[k]) for k in before)  # the click edited nothing
        assert abs(((s.septum_manual[z] - 180) + 180) % 360 - 180) < 2
        wall = panel.last[1]
        assert wall.ed.reference.startswith("septum")
        assert any("by hand" in f for f in wall.flags)
        assert "set by hand" in panel.lbl_auto.text()
        path = s.save_contours(tmp_path / "m.tracecmr.npz")
        s2 = Session(image=s.image, label=np.zeros_like(s.label), spacing=s.spacing)
        s2.load_contours(path)
        assert s2.septum_manual == s.septum_manual
        panel.start_set_septum()
        panel.cancel_click()
        assert panel.editor.click_hook is None
        panel.reset_septum(all_slices=True)
        assert not s.septum_manual and not panel.last[1].ed.reference.startswith("septum")
    finally:
        viewer.close()


def test_reopening_keeps_the_exact_points_no_creeping_tips(tmp_path):
    """Open, edit, save, reopen, many times: a sharp-tipped RV must stay exactly as drawn (rebuilding the points
    from pixels on every open used to round the tips a little each time). Smoothed slices stay smoothed."""
    from tracecmr.app.session import Session as S

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        e = panel.editor
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        epi = e.curves[(t, z, "epi")]
        c = epi.mean(axis=0)
        r = np.hypot(*(epi - c).T).mean()
        ang = np.linspace(200, 340, 9)  # a thin crescent with sharp tips below the LV
        outer = np.column_stack([c[0] + (r + 7) * np.cos(np.deg2rad(ang)), c[1] - (r + 7) * np.sin(np.deg2rad(ang))])
        inner = np.column_stack([c[0] + (r + 2) * np.cos(np.deg2rad(ang[::-1][1:-1])),
                                 c[1] - (r + 2) * np.sin(np.deg2rad(ang[::-1][1:-1]))])
        rv = np.vstack([outer, inner])
        e.push_undo()
        e.curves[(t, z, "rv")] = rv.copy()
        e.edited.add((t, z, "rv"))
        panel.cmb_add.setCurrentIndex(1)
        panel._smooth_points()  # epi smoothed
        path = tmp_path / "m.tracecmr.npz"
        for cycle in range(5):
            panel.save_contours(quiet=True, path=path)
            s2 = S(image=s.image, label=np.zeros_like(s.label), spacing=s.spacing, subject_id="phantom")
            s2.load_contours(path)
            panel.load_session(s2)
            e = panel.editor
            viewer.dims.current_step = (t, z, *viewer.dims.current_step[2:])
            assert np.allclose(e.curves[(t, z, "rv")], rv, atol=1e-4), f"RV moved after reopen {cycle + 1}"
            ek = (t, z, "endo")
            e.move_point(ek, 0, e.curves[ek][0] + 0.01)  # an edit on the slice: it is re-rendered from the points
        # a slice changed outside the editor after saving (e.g. Auto LV) falls back to its pixels
        s3 = S(image=s.image, label=np.zeros_like(s.label), spacing=s.spacing)
        s3.load_contours(path)
        s3.label[:, :, z, t] = 0
        assert not any(k[:2] == (t, z) for k in s3.curves_still_valid())
    finally:
        viewer.close()


def test_export_all_workbook_images_and_study_table(tmp_path, monkeypatch):
    from pathlib import Path

    import openpyxl
    from qtpy.QtCore import QSettings

    from tracecmr.app import report

    store = {"folders/contours": str(tmp_path / "results")}
    monkeypatch.setattr(QSettings, "value", lambda self, key, *a, **k: store.get(key))
    monkeypatch.setattr(QSettings, "setValue", lambda self, key, v: store.__setitem__(key, v))

    def make(name):
        ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
        es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
        return Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                       spacing=ed.spacing, subject_id=name, papillary="blood_pool", ed_frame=0, es_frame=1,
                       source=str(tmp_path / f"{name}.h5"))

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(make("MouseA"))
        out = panel.export_all()
        wb = openpyxl.load_workbook(out["workbook"])
        assert wb.sheetnames == ["Summary", "Per slice", "Wall segments", "Notes"]
        summ = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True) if r[0]}
        assert summ["LV ejection fraction (EF)"] == pytest.approx(panel.last[0].ef_pct, abs=0.01)
        assert summ["Wall thickness method"].startswith("centerline")
        segs = list(wb["Wall segments"].iter_rows(min_row=2, values_only=True))
        assert len(segs) == 16
        imgs = {p.name for p in (out["folder"] / "images").glob("*.png")}
        assert {"bullseye_ED_wall.png", "bullseye_thickening.png", "montage_ED.png"} <= imgs
        study = tmp_path / "results" / "study_summary.xlsx"
        assert Path(out["study"]) == study
        panel.export_all()  # again: the same row is updated, not duplicated
        panel.load_session(make("MouseB"))
        panel.export_all()
        ws = openpyxl.load_workbook(study)["Mice"]
        header = [c.value for c in ws[1]]
        mice = [r[header.index("mouse")] for r in ws.iter_rows(min_row=2, values_only=True)]
        assert sorted(mice) == ["MouseA", "MouseB"]
        assert "ef_pct" in header and "mean_edwt_mm" in header and any(h.startswith("edwt_mm_") for h in header)
        # without openpyxl: the same tables as CSV
        monkeypatch.setattr(report, "_openpyxl", lambda: None)
        out = panel.export_all()
        assert Path(out["workbook"]).is_dir() and (Path(out["workbook"]) / "summary.csv").exists()
        assert Path(out["study"]).suffix == ".csv"
    finally:
        viewer.close()


def test_export_files_are_written_as_utf8(tmp_path, monkeypatch):
    """Notes contain characters like '→'; Windows' default encoding can't write them, so every export file must say
    its encoding (this crashed the first export on Windows)."""
    import builtins
    import pathlib

    from tracecmr.app import report

    real_open, real_write = builtins.open, pathlib.Path.write_text

    def strict_open(file, mode="r", *a, **k):
        if "b" not in mode and "encoding" not in k:
            raise AssertionError(f"{file} opened without an encoding")
        return real_open(file, mode, *a, **k)

    def strict_write(self, data, *a, **k):
        if "encoding" not in k:
            raise AssertionError(f"{self} written without an encoding")
        return real_write(self, data, *a, **k)

    rep = {"summary": [("LV ejection fraction (EF)", 70.0, "%")], "settings": [("Mouse", "M→1")],
           "per_slice": [{"slice": 1, "edv_ul": 1.0}], "segments": [{"aha": 1, "name": "anterior"}],
           "notes": ["tick Settings → Mirror segment numbering", "⚠ check"], "study": {"mouse": "M→1", "ef_pct": 70}}
    monkeypatch.setattr(builtins, "open", strict_open)
    monkeypatch.setattr(pathlib.Path, "write_text", strict_write)
    folder = report._write_mouse_csv(tmp_path / "m", rep)
    monkeypatch.setattr(report, "_openpyxl", lambda: None)
    study = report.update_study(tmp_path / "study_summary.xlsx", rep["study"], key="m")
    study = report.update_study(tmp_path / "study_summary.xlsx", rep["study"], key="m")  # read back + rewrite
    monkeypatch.setattr(builtins, "open", real_open)
    assert "→" in (folder / "notes.txt").read_text(encoding="utf-8")
    assert "M→1" in study.read_text(encoding="utf-8-sig")


def _auto_lv_session(monkeypatch, tweak=None, unsure=None):
    """A phantom mouse contoured by a fake Auto LV. tweak(z, t, lab) may alter the model's output for a slice;
    unsure maps (t, z) to the model-unsure score it reports."""
    from tracecmr import autocontour as ac

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    img = np.stack([ed.image, es.image], -1)
    truth = np.stack([ed.label, es.label], -1)

    def fake_pairs(image, spacing, pairs, model="ucl", unsure_out=None):
        if unsure_out is not None:
            unsure_out[:] = [(unsure or {}).get((t, z), 0.02) for z, t in pairs]
        out = []
        for z, t in pairs:
            lab = truth[:, :, z, t].copy()
            if tweak:
                lab = tweak(z, t, lab)
            out.append(lab)
        return np.stack(out, axis=2)

    monkeypatch.setattr(ac, "available", lambda: {"ucl": "ready"})
    monkeypatch.setattr(ac, "run_pairs", fake_pairs)
    return Session(image=img, label=np.zeros(img.shape, np.uint8), spacing=ed.spacing, subject_id="m",
                   ed_frame=0, es_frame=1), truth


def _contour_both(panel):
    for t in (0, 1):
        panel._show_frame(t)
        panel.autocontour(sync=True)
    panel.compute()


def _flag_texts(panel, z):
    item = next(it for it in panel._items() if it.data(0, Qt.ItemDataRole.UserRole) == z)
    return item.text(0), item.toolTip(0)


def test_double_check_flags_only_on_untouched_model_outlines(monkeypatch):
    """Flags (model unsure, image, anatomy, Auto LV's base/apex notes) appear only while a slice is exactly the
    model's outline; any edit makes it the user's and every flag on it goes."""
    from scipy import ndimage as ndi

    from tracecmr import labels as LB

    def es_bigger(z, t, lab):  # the model draws slice 3's ES cavity bigger than its ED one
        if (z, t) == (3, 1):
            blood = lab == LB.LV_BLOODPOOL
            lab[ndi.binary_dilation(blood, iterations=9) & (lab == LB.LV_MYOCARDIUM)] = LB.LV_BLOODPOOL
        return lab

    s, _ = _auto_lv_session(monkeypatch, tweak=es_bigger, unsure={(0, 4): 0.6})
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(s)
        _contour_both(panel)
        assert any("model unsure" in w for w in panel._checks.get((0, 4), []))
        assert any("ES cavity is bigger" in w for w in panel._checks.get((1, 3), []))
        text, tip = _flag_texts(panel, 4)
        assert text.startswith("⚠") and "Double-check:" in tip and "model unsure" in tip
        assert "Double-check slice(s)" in panel.lbl_notes.text()
        e = panel.editor
        for t, z in ((0, 4), (1, 3)):  # the user edits both slices (a tiny drag is enough)
            step = list(viewer.dims.current_step); step[0], step[1] = t, z; viewer.dims.current_step = tuple(step)
            k = (t, z, "endo")
            e.move_point(k, 0, e.curves[k][0] + 0.05)
        panel.compute()
        assert (0, 4) not in panel._checks and (1, 3) not in panel._checks
        assert "model unsure" not in _flag_texts(panel, 4)[1] and "ES cavity" not in _flag_texts(panel, 3)[1]
    finally:
        viewer.close()


def test_no_flags_on_hand_drawn_slices_or_after_a_brush_edit(monkeypatch):
    """Hand-drawn contours are never second-guessed; a pixel-brush edit (or any label change) also counts as
    the user's outline, including for Auto LV's base/apex notes."""
    from tracecmr import labels as LB

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    lab = np.stack([ed.label, es.label], -1)
    lab[:, :, 3, 1] = lab[:, :, 3, 0]  # hand-drawn ES outline as big as ED: anatomy would object
    s = Session(image=np.stack([ed.image, es.image], -1), label=lab, spacing=ed.spacing, subject_id="h",
                papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(s)
        panel.compute()
        assert panel._checks == {}
        assert not any(_flag_texts(panel, z)[0].startswith("⚠") for z in range(8)
                       if any(it.data(0, Qt.ItemDataRole.UserRole) == z for it in panel._items()))
    finally:
        viewer.close()

    s2, _ = _auto_lv_session(monkeypatch, unsure={(1, 2): 0.6})
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(s2)
        _contour_both(panel)
        assert (1, 2) in panel._checks
        flagged_auto = {z for z in range(8) if any(it.data(0, Qt.ItemDataRole.UserRole) == z for it in panel._items())
                        and "slice:" in _flag_texts(panel, z)[1]}
        lab = panel.session.label.copy()  # a brush edit on every slice: paint one pixel of myocardium
        for z in range(8):
            for t in (0, 1):
                if lab[:, :, z, t].any():
                    yy, xx = np.nonzero(lab[:, :, z, t] == LB.LV_MYOCARDIUM)
                    if len(yy):
                        lab[yy[0], xx[0], z, t] = LB.LV_BLOODPOOL
        viewer.layers["contours"].data = lab.transpose(3, 2, 1, 0).copy()
        panel.rb_pixels.setChecked(True)
        panel.compute()
        assert panel._checks == {}
        for z in flagged_auto:  # Auto LV's base/apex notes are gone too
            assert "slice:" not in _flag_texts(panel, z)[1]
    finally:
        viewer.close()


def test_including_the_papillary_is_not_flagged(monkeypatch):
    """The reported case: the model outline is flagged; the user includes the papillary muscle (Grow / drag),
    and the slice is no longer flagged even though the image check would disagree with the new outline."""
    from tracecmr import labels as LB

    def papillary(z, t, lab):  # the model leaves a big dark notch out of slice 4's ED cavity
        if (z, t) == (4, 0):
            yy, xx = np.nonzero(lab == LB.LV_BLOODPOOL)
            cy, cx = int(np.mean(yy)), int(np.mean(xx))
            lab[cy - 3:cy + 4, cx:cx + 12][lab[cy - 3:cy + 4, cx:cx + 12] == LB.LV_BLOODPOOL] = LB.LV_MYOCARDIUM
        return lab

    s, _ = _auto_lv_session(monkeypatch, tweak=papillary, unsure={(0, 4): 0.3})
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(s)
        _contour_both(panel)
        assert (0, 4) in panel._checks
        step = list(viewer.dims.current_step); step[0], step[1] = 0, 4; viewer.dims.current_step = tuple(step)
        for _ in range(3):
            panel.grow_contour(0.5)
        panel.compute()
        assert (0, 4) not in panel._checks
        assert not _flag_texts(panel, 4)[0].startswith("⚠")
    finally:
        viewer.close()


def test_opening_another_mouse_detaches_the_old_editor(tmp_path):
    """Each opened mouse gets a new editor; the old one must stop listening (it kept handling clicks and slice
    changes on outlines no longer shown), and reopening from a late frame must not loop."""
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)

    def make():
        return Session(image=np.stack([ed.image, es.image, es.image], -1),
                       label=np.stack([ed.label, es.label, np.zeros_like(es.label)], -1),
                       spacing=ed.spacing, subject_id="p", papillary="blood_pool", ed_frame=0, es_frame=1)

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(make())
        old = panel.editor
        viewer.dims.current_step = (2, 7, *viewer.dims.current_step[2:])  # late frame, empty slice
        panel.btn_wall.setChecked(True)
        n_cb = len(viewer.mouse_drag_callbacks)
        panel.load_session(make())
        assert panel.editor is not old
        assert old._on_mouse not in viewer.mouse_drag_callbacks and old.points is None
        assert len(viewer.mouse_drag_callbacks) <= n_cb
        panel.load_session(make())  # and again: no pile-up, no loop
        assert len(viewer.mouse_drag_callbacks) <= n_cb
        assert panel.editor.points is not None and "control points" in viewer.layers
    finally:
        viewer.close()


def test_app_icon_is_the_tracecmr_logo():
    from tracecmr.app.launch import ICON, set_icon

    assert ICON.exists() and ICON.with_suffix(".ico").exists()
    viewer = napari.Viewer(show=False)
    try:
        set_icon(viewer)
        icon = viewer.window._qt_window.windowIcon()
        assert not icon.isNull() and icon.availableSizes()
    finally:
        viewer.close()


def test_start_screen_is_tracecmr_not_napari(tmp_path):
    from qtpy.QtCore import QMimeData, QUrl
    from qtpy.QtWidgets import QLabel

    from tracecmr import __version__
    from tracecmr.app import branding

    assert branding.ICON_SVG.exists() and all(p.exists() for p in branding.LOGO.values())
    viewer = napari.Viewer(show=False)
    opened = []
    try:
        assert branding.start_screen(viewer, open_path=opened.append)
        w = viewer.window._qt_viewer._welcome_widget
        # napari's logo, "drag file(s) here", shortcuts and "Did you know?" tips are hidden
        for part in (w._image, w._label, w._version_label, w._tip_label):
            assert not part.isVisibleTo(w)
        screen = w._tracecmr
        assert screen.isVisibleTo(w)
        texts = " ".join(l.text() for l in screen.findChildren(QLabel))
        assert "Open a scan" in texts and __version__ in texts
        assert "napari" not in texts.lower() and "did you know" not in texts.lower()
        logo = screen.findChild(QLabel, "tracecmr_logo")
        assert not logo.pixmap().isNull()
        viewer.theme = "light"  # logo and text recoloured for the light theme without errors
        assert not logo.pixmap().isNull()

        # a mouse folder dropped on the start screen opens in TraceCMR, not napari's reader
        class Ev:
            def __init__(self, md):
                self.md = md

            def mimeData(self):
                return self.md

            def acceptProposedAction(self):
                pass

        md = QMimeData()
        md.setUrls([QUrl.fromLocalFile(str(tmp_path))])
        w.urls_dropped.emit(Ev(md))
        assert opened == [tmp_path]
    finally:
        viewer.close()


def test_interface_font_is_bundled_inter():
    from qtpy.QtGui import QFontInfo
    from qtpy.QtWidgets import QApplication

    from tracecmr.app import branding

    files = sorted(p.name for p in branding.FONT_DIR.glob("*.ttf"))
    assert files == ["Inter-Bold.ttf", "Inter-Regular.ttf"] and (branding.FONT_DIR / "OFL.txt").exists()
    assert "Inter" in branding.load_fonts()
    app = QApplication.instance()
    before = app.font()
    try:
        assert branding.use_app_font()
        assert QFontInfo(app.font()).family() == "Inter"
    finally:
        app.setFont(before)
    for logo in branding.LOGO.values():  # wordmark is outlines, so it never depends on an installed font
        assert "<text" not in logo.read_text()


def test_splash_screen_shows_and_closes():
    from tracecmr.app.branding import Splash

    s = Splash()
    try:
        assert s.widget.isVisible() and s.text == "Starting" and s.fraction == 0
        s.progress(0.5)
        s.progress(0.3, "Almost ready")  # the line never goes backwards
        assert s.fraction == 0.5 and s.text == "Almost ready"
    finally:
        s.close()
    assert not s.widget.isVisible()


def test_desktop_app_writes_messages_to_a_log(tmp_path, monkeypatch):
    import sys

    from tracecmr.app import launch

    log = tmp_path / "logs" / "tracecmr_log.txt"
    monkeypatch.setenv("TRACECMR_LOG", str(log))
    monkeypatch.setattr(sys, "stdout", None)  # what pythonw / tracecmr-app gives the program
    monkeypatch.setattr(sys, "stderr", None)
    import faulthandler

    try:
        assert launch._windowless_output() == log
        print("hello from the app")
        sys.stdout.flush()
    finally:
        faulthandler.disable()
    text = log.read_text(encoding="utf-8")
    assert "TraceCMR started" in text and "hello from the app" in text


def test_model_folder_from_environment_is_remembered(qtbot, monkeypatch, tmp_path):
    from tracecmr.app._widget import _settings

    st = _settings()
    old = st.value("models/ucl")
    st.remove("models/ucl")
    monkeypatch.setenv("TRACECMR_UCL_DIR", str(tmp_path))
    viewer = napari.Viewer(show=False)
    try:
        TraceCMRWidget(viewer)
        assert st.value("models/ucl") == str(tmp_path)
    finally:
        viewer.close()
        if old is not None:
            st.setValue("models/ucl", old)
        else:
            st.remove("models/ucl")


def test_charts_render_once_after_resizing_stops(qtbot):
    from matplotlib.figure import Figure

    from tracecmr.app._widget import SmoothCanvas

    fig = Figure(figsize=(3, 2))
    fig.add_subplot().plot([0, 1], [0, 1])
    c = SmoothCanvas(fig, delay_ms=50)
    qtbot.addWidget(c)
    c.resize(300, 200)
    c.show()
    qtbot.waitExposed(c)
    c.draw()
    renders = []
    c.mpl_connect("resize_event", lambda e: renders.append(e))
    for w in range(310, 400, 10):  # a drag: many sizes in quick succession
        c.resize(w, 200)
        qtbot.wait(5)
    assert renders == [] and c._pending is not None  # nothing re-rendered while dragging
    qtbot.wait(150)
    assert len(renders) == 1 and c._pending is None  # one crisp render at the final size
    assert c.get_width_height()[0] == c.width()


def test_slice_table_headers_fit():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(s)
        panel.compute()
        hdr = panel.tbl_slices.horizontalHeader()
        fm = hdr.fontMetrics()
        panel.tbl_slices.resize(260, 300)  # a narrow side panel
        for c in range(1, 5):
            text = panel.tbl_slices.horizontalHeaderItem(c).text()
            assert "\n" not in text and hdr.sectionSize(c) >= fm.horizontalAdvance(text)
        assert sum(hdr.sectionSize(c) for c in range(5)) <= panel.tbl_slices.viewport().width() + 1
    finally:
        viewer.close()


def test_panels_can_be_hidden_and_brought_back():
    from qtpy.QtWidgets import QDockWidget

    from tracecmr.app._widget import setup_dock
    from tracecmr.app.launch import add_panels_menu

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        right = viewer.window.add_dock_widget(panel, name="TraceCMR", area="right")
        setup_dock(right)
        docks = [right] + panel.docks
        assert sorted(d.windowTitle() for d in docks) == ["Results", "Scans", "Time bar", "Toolbar"]
        for d in docks:
            assert d.features() == QDockWidget.DockWidgetFeature.DockWidgetClosable  # hide yes, tear off no
        assert hasattr(right.titleBarWidget(), "close_btn")  # side panels: header with a ×
        add_panels_menu(viewer, docks)
        menu = viewer.window._tracecmr_panels_menu
        names = [a.text() for a in menu.actions() if a.text()]
        assert names == ["Scans", "Toolbar", "Results", "Time bar", "Show all panels"]
        viewer.window._qt_window.show()
        right.titleBarWidget().close_btn.click()  # the ×
        assert right.isHidden() and "Results button" in viewer.status
        tick = next(a for a in menu.actions() if a.text() == "Results")
        assert not tick.isChecked()
        tick.trigger()  # tick it again in the menu
        assert not right.isHidden()
        for d in docks:
            d.close()
        viewer.window._tracecmr_restore.trigger()  # Ctrl+Shift+R
        assert all(not d.isHidden() for d in docks)
    finally:
        viewer.close()


def test_logo_trace_draw_in():
    from tracecmr.app import branding

    full = branding._logo_pixmap("dark", 280, 1.0).toImage()
    none = branding._logo_pixmap("dark", 280, 0.0).toImage()
    half = branding._logo_pixmap("dark", 280, 0.5).toImage()

    def red(img):
        n = 0
        for x in range(0, img.width(), 2):
            for y in range(0, img.height(), 2):
                c = img.pixelColor(x, y)
                n += c.alpha() > 100 and c.red() > 150 and c.green() < 140
        return n

    assert red(none) == 0 < red(half) < red(full)


def test_scans_results_focus_buttons(qtbot):
    from tracecmr.app._widget import setup_dock

    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    s = Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        right = viewer.window.add_dock_widget(panel, name="TraceCMR", area="right")
        setup_dock(right)
        win = viewer.window._qt_window
        win.resize(1400, 900)
        viewer.window.show()
        panel.load_session(s)
        qtbot.waitUntil(lambda: getattr(panel, "_panels_wired", False), timeout=3000)
        left, r = panel._side_docks()
        assert r is right and left is not None
        assert panel.btn_left.isChecked() and panel.btn_right.isChecked() and not panel.btn_focus.isChecked()

        panel.btn_right.click()  # hide Results with its button
        assert right.isHidden() and not panel.btn_right.isChecked()
        right.titleBarWidget().close_btn.click() if not right.isHidden() else None
        panel.btn_right.click()  # and back
        assert not right.isHidden()
        left.titleBarWidget().close_btn.click()  # hide Scans with its ×: the button follows
        assert not panel.btn_left.isChecked()
        panel.btn_left.click()
        assert not left.isHidden()

        z0 = panel._camera.zoom
        panel.btn_focus.click()  # Focus: both side panels go, the image grows to use the room
        qtbot.wait(100)
        time_bar = next(d for d in panel.docks if d.objectName() == "Time")
        assert left.isHidden() and right.isHidden() and time_bar.isHidden() and panel.btn_focus.isChecked()
        assert panel._camera.zoom > z0 * 1.05
        panel.btn_focus.click()  # back to the layout before
        qtbot.wait(100)
        assert not left.isHidden() and not right.isHidden() and not time_bar.isHidden()
        assert abs(panel._camera.zoom / z0 - 1) < 0.02  # and the same framing as before
    finally:
        viewer.close()


def _phantom_session():
    ed = tapered_phantom(n_slices=8, base_blood_mm=2.4, apex_blood_mm=1.0, wall_mm=1.0)
    es = tapered_phantom(n_slices=8, base_blood_mm=1.6, apex_blood_mm=0.6, wall_mm=1.3)
    return Session(image=np.stack([ed.image, es.image], -1), label=np.stack([ed.label, es.label], -1),
                   spacing=ed.spacing, subject_id="phantom", papillary="blood_pool", ed_frame=0, es_frame=1)


def test_toolbar_groups_and_dropdowns(qtbot):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        shown = [b for grp in panel.toolbar_groups for b in grp]
        assert len(shown) <= 18  # one button per job
        for b in (panel.btn_ucl, panel.btn_snap, panel.btn_fit, panel.btn_grow, panel.btn_shrink, panel.btn_raw,
                  panel.btn_contrast, panel.btn_zoom_in, panel.btn_zoom_out):
            assert b not in shown and b.isHidden()
        texts = [a.text() for a in panel.btn_refine.menu().actions()]
        assert any(t.startswith("Grow") for t in texts) and any(t.startswith("Snap") for t in texts)
        assert [a.text() for a in panel.btn_dcm.menu().actions()][1].startswith("Add UCL")
        # a menu entry does what its folded-away button does
        hits = []
        panel.btn_grow.clicked.connect(lambda: hits.append(1))
        next(a for a in panel.btn_refine.menu().actions() if a.text().startswith("Grow")).trigger()
        assert hits
        # universal keys: + / − zoom (also with Ctrl and on the keypad), Ctrl+0 centers, Delete clears
        keys = set(getattr(panel, "_key_shortcuts", []))
        assert {"+", "=", "-", "Control-0", "Delete", "Backspace"} <= keys
        # every dropdown entry is a styled row with its keycaps
        grow = next(a for a in panel.btn_refine.menu().actions() if a.text().startswith("Grow"))
        assert [c.text() for c in grow.defaultWidget().caps] == ["]"]
        # Hide contours in View follows the H key / Raw button both ways
        panel.act_raw.setChecked(True)
        assert panel.btn_raw.isChecked()
        panel.btn_raw.setChecked(False)
        assert not panel.act_raw.isChecked()
    finally:
        viewer.close()


def test_smooth_image_and_structure_keys(qtbot):
    from tracecmr.app._widget import _settings

    before = _settings().value("view/smooth_image")
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        panel.set_smooth_image(True)
        assert viewer.layers["image"].interpolation2d == "cubic" and panel.act_smooth_image.isChecked()
        panel.set_smooth_image(False)
        assert viewer.layers["image"].interpolation2d == "nearest" and not panel.act_smooth_image.isChecked()
        panel.btn_endo.click()
        panel.cycle_structure()
        assert panel.btn_epi.isChecked()
        panel.cycle_structure()
        panel.cycle_structure()
        assert panel.btn_endo.isChecked()  # RV wraps back to Endo
    finally:
        viewer.close()
        if before is None:
            _settings().remove("view/smooth_image")
        else:
            _settings().setValue("view/smooth_image", before)


def test_shortcut_sheet_lists_the_bound_keys(qtbot):
    from tracecmr.app.help import SHORTCUTS, ShortcutSheet

    keys = {k for _, rows in SHORTCUTS for k, _ in rows}
    for k in ("Q", "G", "K", "H", "W", "F", "?", "Ctrl+Z", "Ctrl+S"):
        assert k in keys
    sheet = ShortcutSheet()
    qtbot.addWidget(sheet)
    sheet.popup()
    assert sheet.isVisible()
    qtbot.keyClick(sheet, Qt.Key.Key_Escape)
    assert not sheet.isVisible()


def test_first_run_tour(qtbot, monkeypatch):
    from tracecmr.app._widget import _settings, setup_dock

    before = _settings().value("tour/seen")
    monkeypatch.delenv("TRACECMR_NO_TOUR", raising=False)
    _settings().remove("tour/seen")
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        right = viewer.window.add_dock_widget(panel, name="TraceCMR", area="right")
        setup_dock(right)
        win = viewer.window._qt_window
        win.resize(1400, 900)
        viewer.window.show()
        panel.load_session(_phantom_session())
        qtbot.waitUntil(lambda: getattr(panel, "_panels_wired", False), timeout=3000)
        panel.maybe_start_tour()
        tour = panel._tour
        assert tour.isVisible() and tour.index == 0 and tour.geometry() == win.rect()
        for i in range(1, len(tour.steps)):
            qtbot.keyClick(tour, Qt.Key.Key_Right)
            assert tour.index == i
            if tour.steps[i].target is not None:
                assert not tour._target_rect(tour.steps[i]).isNull(), tour.steps[i].title  # each stop finds its part
        qtbot.keyClick(tour, Qt.Key.Key_Left)
        assert tour.index == len(tour.steps) - 2
        tour.next()
        tour.next()  # Done
        assert not tour.isVisible()
        assert str(_settings().value("tour/seen")).lower() == "true"
        panel._tour = None
        panel.maybe_start_tour()  # seen: not again
        assert panel._tour is None
    finally:
        viewer.close()
        if before is None:
            _settings().remove("tour/seen")
        else:
            _settings().setValue("tour/seen", before)


def test_roi_tool_measures_compares_and_exports(qtbot, tmp_path):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        assert not panel.roi.table.isVisible() or panel.roi.table.rowCount() == 0
        panel.session.image = panel.session.image + np.random.default_rng(0).normal(0, 0.05, panel.session.image.shape)
        panel.btn_roi.setChecked(True)
        assert panel.roi.active and panel.editor.suspended
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        panel.roi.add({"kind": "circle", "t": t, "z": z, "cx": 30.0, "cy": 30.0, "r_mm": 0.3})
        panel.roi.add({"kind": "freehand", "t": t, "z": z, "pts": [[2, 2], [8, 2], [8, 8], [2, 8]]})
        assert panel.roi.table.rowCount() == 2 and "rois" in viewer.layers
        assert len(viewer.layers["rois"].data) == 2  # both drawn on this slice
        r = panel.roi.ratios()
        assert r is not None and np.isfinite(r["snr_a"])
        # rename by editing the name cell: the comparison follows
        panel.roi.table.item(0, 0).setText("blood")
        assert panel.session.rois[0]["name"] == "blood" and panel.roi.cmb_a.currentText() == "blood"
        out = panel.roi.export_csv(tmp_path / "rois.csv")
        text = out.read_text()
        assert "blood" in text and "CNR" in text
        # Delete in ROI mode removes the selected region, not the contour
        panel.roi.table.selectRow(1)
        panel._delete_key()
        assert len(panel.session.rois) == 1
        panel._escape()  # Esc leaves ROI mode and editing comes back
        assert not panel.btn_roi.isChecked() and not panel.editor.suspended
    finally:
        viewer.close()


def test_hide_contours_comes_back_without_rebuilding(qtbot):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        e = panel.editor
        z = int(viewer.dims.current_step[1])
        e.move_point((0, z, "endo"), 0, e.curves[(0, z, "endo")][0] + 0.7)  # an edit not yet written back
        curves_before = e.curves
        panel.btn_raw.setChecked(True)
        assert not viewer.layers["curves"].visible and not viewer.layers["control points"].visible
        panel.btn_raw.setChecked(False)
        assert viewer.layers["curves"].visible and viewer.layers["control points"].visible
        assert e.curves is curves_before and (0, z, "endo") in e.edited  # same curves, nothing rebuilt or lost
    finally:
        viewer.close()


def test_toolbar_hover_cards():
    from tracecmr.app.hints import PROP, HintCard

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        for grp in panel.toolbar_groups:
            for b in grp:
                hint = b.property(PROP)
                assert hint and hint["title"] and (hint["body"] or hint["keys"]), b.text()
        assert panel.btn_smooth.property(PROP)["keys"] == ["S"]
        card = HintCard()
        card.fill(panel.btn_save.property(PROP))
        assert card.title.text() == "Save" and card.keys_row.count() == 2
        assert {"s", "p"} <= set(panel._key_shortcuts)
    finally:
        viewer.close()


def test_letter_keys_work_after_clicking_a_list_or_button(qtbot):
    """S, P and + reach TraceCMR wherever the focus is, not only when the image has it."""
    from qtpy.QtTest import QTest

    from tracecmr.app._widget import setup_dock

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        setup_dock(viewer.window.add_dock_widget(panel, name="TraceCMR", area="right"))
        viewer.window.show()
        panel.load_session(_phantom_session())
        win = viewer.window._qt_window
        calls = []
        orig = panel._smooth_points
        panel._smooth_points = lambda: (calls.append(1), orig())
        assert panel.list.focusPolicy() == Qt.FocusPolicy.NoFocus  # the slice list never takes the keyboard
        panel.btn_export.setFocus()  # e.g. after clicking a button in the Results panel
        QTest.keyClick(QApplication.focusWidget() or win, Qt.Key.Key_S)
        assert calls
        QTest.keyClick(QApplication.focusWidget() or win, Qt.Key.Key_P)
        assert not panel.btn_points.isChecked()
        z0 = panel._camera.zoom
        QTest.keyClick(QApplication.focusWidget() or win, Qt.Key.Key_Plus, Qt.KeyboardModifier.ShiftModifier, )
        assert panel._camera.zoom > z0
    finally:
        viewer.close()


def test_results_tabs_follow_the_tools(qtbot):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        tabs = panel.results_tabs
        assert list(tabs.index) == ["function", "wall", "regions"] and tabs.current() == "function"
        panel.btn_roi.setChecked(True)  # the ROI tool opens the Regions tab
        assert tabs.current() == "regions"
        panel.btn_roi.setChecked(False)
        panel.btn_wall.setChecked(True)  # Wall opens the Wall tab, and its tick there follows
        assert tabs.current() == "wall" and panel.chk_wall_lines.isChecked()
        panel.rb_wall_thickening.setChecked(True)
        assert panel.wall_mode == "thickening"
        tabs.show("function")
        assert panel.results_stack.currentIndex() == 0 and tabs.buttons["function"].isChecked()
        tabs.show("wall")  # picking Wall turns the lines on
        panel.chk_wall_lines.setChecked(False)
        assert not panel.btn_wall.isChecked()
        assert {"Control-1", "Control-2", "Control-3"} <= set(panel._key_shortcuts)
    finally:
        viewer.close()


def test_results_update_without_a_compute_button(qtbot):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        assert panel.btn_compute.isHidden()
        qtbot.waitUntil(lambda: panel.res_labels["ef_pct"].text() not in ("–", ""), timeout=3000)
        e = panel.editor
        z = int(viewer.dims.current_step[1])
        before = panel.res_labels["edv_ul"].text()
        e.offset(0, z, 1.0)  # grow the endocardium: EDV follows by itself
        qtbot.waitUntil(lambda: panel.res_labels["edv_ul"].text() != before, timeout=3000)
    finally:
        viewer.close()


def test_redo_puts_back_an_undone_edit(qtbot):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        e = panel.editor
        z = int(viewer.dims.current_step[1])
        key = (0, z, "endo")
        before = e.curves[key].copy()
        e.offset(0, z, 1.0)
        grown = e.curves[key].copy()
        panel.undo()
        assert np.allclose(e.curves[key], before)
        panel.redo()
        assert np.allclose(e.curves[key], grown)
        panel.redo()  # nothing more to redo
        assert "Nothing to redo" in panel.lbl_auto.text()
        panel.undo()
        e.offset(0, z, -0.5)  # a new edit ends the redo history
        assert not e._redo
    finally:
        viewer.close()


def test_autosave_and_recovery(qtbot, tmp_path, monkeypatch):
    from qtpy.QtWidgets import QMessageBox

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        monkeypatch.setattr(panel, "contours_dir", lambda: tmp_path)
        s = _phantom_session()
        s.source = str(tmp_path / "MouseX")
        panel.load_session(s)
        qtbot.wait(20)
        assert panel.autosave.tick() is None  # nothing changed yet: nothing written
        z = int(viewer.dims.current_step[1])
        panel.editor.offset(0, z, 1.0)
        rec = panel.autosave.tick()
        assert rec is not None and rec.exists()
        changed = panel.editor.apply_to(panel.session.label)
        # reopen the same scan as if TraceCMR had crashed: the change is offered and restored
        s2 = _phantom_session()
        s2.source = s.source
        panel._recovery_answer = QMessageBox.StandardButton.Yes
        panel.load_session(s2)
        qtbot.waitUntil(lambda: "restored" in panel.lbl_auto.text(), timeout=3000)
        assert np.array_equal(panel.session.label, changed)
        panel._sync_labels()  # the contours on screen are the restored ones too (not the pre-crash points)
        assert np.array_equal(panel.session.label, changed)
        assert np.array_equal(panel.editor.apply_to(panel.session.label), changed)
        panel.save_contours(quiet=True, path=tmp_path / "MouseX.tracecmr.npz")
        assert not rec.exists()  # a normal save removes the recovery file
    finally:
        viewer.close()


def test_ruler(qtbot):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        dx, dy, _ = panel.session.spacing
        panel.roi.add({"kind": "line", "t": t, "z": z, "pts": [[10.0, 10.0], [10.0 + 3 / dx, 10.0 + 4 / dy]]})
        assert panel.session.rois[-1]["name"] == "Ruler 1"
        assert panel.roi.table.item(0, 1).text() == "5.00 mm"
        assert "5.00 mm" in panel.lbl_auto.text()
        assert panel.roi.ratios() is None  # a ruler is not a region to compare
    finally:
        viewer.close()


def test_problem_report(qtbot, tmp_path, monkeypatch):
    monkeypatch.setenv("TRACECMR_LOG", str(tmp_path / "logs" / "tracecmr_log.txt"))
    (tmp_path / "logs").mkdir()
    (tmp_path / "logs" / "tracecmr_log.txt").write_text("hello\nsomething went wrong\n")
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel._report_problem()
        reports = list((tmp_path / "logs").glob("TraceCMR_problem_report_*.txt"))
        assert reports and "something went wrong" in reports[0].read_text() and "TraceCMR" in reports[0].read_text()
    finally:
        viewer.close()


def test_ed_es_side_by_side_and_figures(qtbot, tmp_path):
    from tracecmr.app.compare import save_all_slices_figure, save_slice_figure, slice_contours

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        panel.show_compare()
        win = panel._compare
        assert win.is_open() and "slice" in win.title.text()
        assert win.dock is None  # a view over the image, not a panel or a popup
        panel.show_compare()  # Ctrl+E again hides it
        assert not win.is_open()
        panel.show_compare()
        assert win.is_open()
        z = panel.current_slice()
        assert {"endo", "epi"} <= set(slice_contours(panel, panel.session.t_for(z, "ES"), z))
        panel.step_slice(1)  # follows the slice
        qtbot.waitUntil(lambda: f"slice {panel.current_slice()}" in win.title.text(), timeout=2000)
        one = save_slice_figure(panel, panel.current_slice(), tmp_path / "pair.png", **win.opts())
        allz = save_all_slices_figure(panel, tmp_path / "all.pdf", **win.opts())
        assert one.stat().st_size > 10_000 and allz.stat().st_size > 5_000
        from PIL import Image

        w, h = Image.open(one).size
        assert w >= 1500  # 300 dpi
    finally:
        viewer.close()


def test_fast_dot_drag_keeps_layers_in_sync():
    import numpy as np
    from vispy.visuals.mesh import MeshVisual

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        e = panel.editor
        t, z = e._here()
        key = next(k for k in sorted(e.curves) if k[:2] == (t, z))
        e._last_draw = 0.0
        e.move_point(key, 2, e.curves[key][2] + 1.5)
        e._drag_update(key, 2)  # only that dot and that contour are redrawn
        row = [r for r, k in enumerate(e._keys) if k == key][2]
        x, y = e.curves[key][2]
        assert np.allclose(e.points.data[row, 2:], (y, x))
        assert np.allclose(e.shapes.data[e._shape_of[key]], e._outline(key))
        mesh = MeshVisual()  # shapes redraws reuse their colour function, so the shader is not recompiled
        rgba = np.ones((3, 4), np.float32)
        assert mesh._build_color_transform(rgba) is mesh._build_color_transform(rgba)
    finally:
        viewer.close()


def test_dot_grab_radius_grows_with_zoom():
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        e = panel.editor
        viewer.camera.zoom = 10
        far = e._tol_px()
        viewer.camera.zoom = 2000  # zoomed right in: the drawn dot is bigger than 12 screen pixels
        assert e._tol_px() >= e.DOT_SIZE / 2 + 0.25 and e._tol_px() <= far
        t, z = e._here()
        key = next(k for k in sorted(e.curves) if k[:2] == (t, z))
        p = e.curves[key][0] + (0.6, 0.3)  # on the dot's rim (0.67 px from its centre)
        assert e.hit(t, z, p, e._tol_px()) == ("point", key, 0)
    finally:
        viewer.close()


def test_ed_es_view_zooms_and_pans_together(qtbot):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        panel.show_compare()
        v = panel._compare.view
        assert v.zoom == 1.0
        panel.zoom(2.0)  # + goes to the pair while the view is up
        assert v.zoom == 2.0
        v.pan = [30.0, -10.0]
        panel.fit_heart()  # F / Ctrl+0 fit the pair again
        assert v.zoom == 1.0 and v.pan == [0.0, 0.0]
        v.zoom_by(0.5)  # never below fit
        assert v.zoom == 1.0
    finally:
        viewer.close()


def test_ed_es_same_crop_for_all_slices(qtbot, tmp_path):
    from tracecmr.app.compare import pair_frames, save_all_slices_figure

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        s = panel.session
        zs = [z for z in s.order if z in s.included["ED"]]
        per = {z: pair_frames(panel, z, crop=True)["box"] for z in zs}
        same = {z: pair_frames(panel, z, crop="all")["box"] for z in zs}
        assert len(set(same.values())) == 1  # one box for every slice
        box = next(iter(same.values()))
        for z, (xa, xb, ya, yb) in per.items():  # ... that contains each slice's own crop
            assert box[0] <= xa and box[1] >= xb and box[2] <= ya and box[3] >= yb
        panel.show_compare()
        panel._compare.cmb_crop.setCurrentIndex(1)
        assert panel._compare.opts()["crop"] == "all"
        out = save_all_slices_figure(panel, tmp_path / "all.png", **panel._compare.opts())
        assert out.stat().st_size > 5_000
    finally:
        viewer.close()


def test_ed_es_lock_crop_keeps_one_window_on_every_slice(qtbot):
    from tracecmr.app.compare import pair_frames

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        panel.show_compare()
        c = panel._compare
        c.view.resize(800, 400)
        c.refresh(now=True)
        c.view.zoom_by(2.0)  # zoom in on this slice, then lock
        c.btn_lock.setChecked(True)
        box = c.locked_box
        assert box is not None and c.view.zoom == 1.0 and c.opts()["crop"] == box
        own = pair_frames(panel, panel.current_slice(), crop=True)["box"]
        assert box[1] - box[0] < own[1] - own[0]  # the zoomed window is smaller than the slice's own crop
        panel.step_slice(1)
        qtbot.waitUntil(lambda: c.view.frames["box"] == box, timeout=2000)  # the same window on the next slice
        c.btn_lock.setChecked(False)
        assert c.opts()["crop"] is True
    finally:
        viewer.close()


def test_wall_tab_is_ready_before_it_is_opened(qtbot):
    """The slow parts of the Wall tab happen while the mouse is opening, not on the first click."""
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        panel._prewarm_wall()  # normally a moment after loading
        assert "wall thickness" in viewer.layers and not viewer.layers["wall thickness"].visible
        panel.compute()
        drawn = panel._bullseye_drawn
        assert drawn is not None
        panel._draw_bullseyes()  # same numbers, same size: nothing is rebuilt
        assert panel._bullseye_drawn is drawn
        axes = list(panel.fig.axes)
        panel.results_tabs.show("wall")
        assert list(panel.fig.axes) == axes  # picking the tab shows what is already drawn
        panel._bullseye_wall = None  # new numbers: it does redraw
        panel._draw_bullseyes()
        assert panel._bullseye_drawn is None
    finally:
        viewer.close()


def test_detect_reports_ed_es_timing_per_slice(qtbot, tmp_path):
    from tests.test_drift import _cine

    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        s = _cine(drift_frames=4)
        s.separate_slices = True
        panel.load_session(s)
        panel._auto_frames()
        text = panel.lbl_order.text()
        assert "ES varies by" in text and "per slice" in text
    finally:
        viewer.close()


def test_rotate_and_mirror_the_picture(qtbot, tmp_path):
    """Turning the picture is display only: every layer turns together, clicks still land on the right pixel,
    ED | ES and figures turn the same way, and the choice is saved with the mouse."""
    from tracecmr.app.compare import save_slice_figure
    from tracecmr.app.export_images import export_images

    viewer = napari.Viewer(show=False)
    panel = None
    try:
        panel = TraceCMRWidget(viewer)
        panel.load_session(_phantom_session())
        before = panel.session.label.copy()
        panel.set_orientation(angle=90, flip=True)
        aff = panel._orient_affine()
        for layer in viewer.layers:
            if layer.ndim == 4:
                assert np.allclose(np.asarray(layer.affine.affine_matrix), aff), layer.name
        assert np.array_equal(panel.session.label, before)  # nothing about the data changed
        # drag a control point in the turned view: it moves to the pixel under the cursor
        e = panel.editor
        t, z = (int(v) for v in viewer.dims.current_step[:2])
        key = (t, z, "endo")
        p0 = e.curves[key][0].copy()
        c = e.curves[key].mean(0)
        target = c + 1.3 * (p0 - c)
        img = viewer.layers["image"]
        world = lambda p: tuple(img.data_to_world((t, z, p[1], p[0])))  # noqa: E731
        _drive(e, viewer, _Ev(world(p0)), [(world(target), (5, 5))])
        assert np.allclose(e.curves[key][0], target, atol=1e-6)
        # a layer added later (the wall lines) turns too
        panel.btn_wall.setChecked(True)
        qtbot.waitUntil(lambda: "wall thickness" in viewer.layers, timeout=3000)
        assert np.allclose(np.asarray(viewer.layers["wall thickness"].affine.affine_matrix), aff)
        # ED | ES turns the same way and still saves figures; the popover controls follow
        panel.show_compare()
        win = panel._compare
        qtbot.waitUntil(lambda: win.view.frames is not None and win.view.frames["angle"] == 90, timeout=2000)
        assert win.view.frames["flip"] is True and win.view.visible_box() is not None
        out = save_slice_figure(panel, panel.current_slice(), tmp_path / "turned.png", **win.opts())
        assert out.stat().st_size > 10_000
        panel.show_compare()
        panel.show_orientation()
        assert panel.sld_rotation.value() == 90 and panel.btn_mirror.isChecked()
        panel.rotate_by(-90)
        assert panel.orientation() == (0.0, True)
        assert export_images(panel.session, tmp_path / "img")
        # saved with the mouse, and the next mouse starts the same way
        panel.set_orientation(angle=-45, flip=False)
        path = panel.session.save_contours(tmp_path / "m.tracecmr.npz")
        s2 = _phantom_session()
        s2.load_contours(path)
        assert s2.view_rotation == -45 and s2.view_flip is False
        s3 = _phantom_session()
        panel.load_session(s3)
        assert panel.orientation() == (-45.0, False)
    finally:
        if panel is not None and panel.session is not None:
            panel.set_orientation(angle=0, flip=False)  # leave the remembered default as it was
        viewer.close()


def test_auto_lv_offers_one_click_setup_when_not_installed(qtbot, tmp_path, monkeypatch):
    """Clicking Auto LV without the model offers the setup; after it, the model folder is set and Auto LV runs."""
    from tracecmr import autocontour as ac
    from tracecmr import model_setup as ms

    monkeypatch.setenv("TRACECMR_UCL_DIR", str(tmp_path / "missing"))
    monkeypatch.setattr(ms, "default_dir", lambda: tmp_path / "ucl")
    present = {"tf": False}
    monkeypatch.setattr(ms, "runtime_present", lambda: present["tf"])
    monkeypatch.setattr(ac, "_present", lambda m: present["tf"])
    monkeypatch.setattr(ac.sys, "prefix", str(tmp_path / "prefix"))  # no model shipped with this test Python

    def fake_setup(dest, cancelled=lambda: False):
        yield ms.Step("model", 50, 100, "Downloading the model: model weights (156 MB)")
        dest.mkdir(parents=True)
        (dest / "ucl_seg459.onnx").write_bytes(b"0" * 2_000_000)
        yield ms.Step("runtime", 0, 0, "Installing ONNX Runtime (about 20 MB)")
        present["tf"] = True
        yield ms.Step("done", 1, 1, "Auto-contouring is ready")

    monkeypatch.setattr(ms, "run_setup", fake_setup)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(_phantom_session())
        assert not panel.model_ready()
        assert "one-time download" in panel.lbl_auto.text()
        before = panel.session.label.copy()
        panel._setup_answer = False  # "Not now": nothing changes
        panel.autocontour(sync=True)
        assert np.array_equal(panel.session.label, before) and not panel.model_ready()
        panel._setup_answer = True  # "Set up now"
        panel.autocontour(sync=True)
        assert panel.model_ready() and ac.available()["ucl"] == "ready"
        assert panel.setup_dialog.lbl_step.text() == "Auto-contouring is ready"
        assert "ready" in panel.lbl_auto.text()
        assert not panel.btn_setup_model.isEnabled()
        # a failed setup explains itself and offers to try again
        present["tf"] = False

        def broken(dest, cancelled=lambda: False):
            yield ms.Step("runtime", 0, 0, "Installing ONNX Runtime (about 20 MB)")
            raise ms.SetupError("Installing TensorFlow failed. Check the internet connection and try again.")

        monkeypatch.setattr(ms, "run_setup", broken)
        panel.start_model_setup(sync=True)
        assert panel.setup_dialog.lbl_step.text() == "Setup stopped"
        assert "internet" in panel.setup_dialog.lbl_text.text() and panel.setup_dialog.btn_retry.isVisibleTo(panel.setup_dialog)
    finally:
        ac.reset_workers()
        viewer.close()


def test_update_card_flow(qtbot, tmp_path, monkeypatch):
    """A newer version: the card offers it, downloads it, and installs on restart or at the next close. Up to
    date, offline, needing the installer, and the note after an update each say so plainly."""
    from tracecmr import updater as up

    monkeypatch.setattr(up, "updates_dir", lambda: tmp_path / "updates")
    monkeypatch.setattr(up, "log_path", lambda: tmp_path / "update_log.txt")
    monkeypatch.setattr(up, "install_kind", lambda: "app")
    up.cancel_scheduled()
    new = up.Update(version="9.9.9", notes=["Rotate the picture. More detail here.", "Faster."],
                    installer_url="https://example.org/TraceCMR-macOS.pkg")
    answer = {"check": new}

    def fake_check(url=None):
        if isinstance(answer["check"], Exception):
            raise answer["check"]
        return answer["check"]

    monkeypatch.setattr(up, "check", fake_check)
    wheel = tmp_path / "updates" / "tracecmr-9.9.9-py3-none-any.whl"

    def fake_download(u, progress=None, cancelled=lambda: False):
        wheel.parent.mkdir(parents=True, exist_ok=True)
        wheel.write_bytes(b"w")
        progress and progress(1, 1)
        return wheel

    monkeypatch.setattr(up, "download", fake_download)
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        ctl = panel.updates
        ctl.app_mode = True
        ctl.check(manual=False, sync=True)
        card = ctl.card
        assert card.state == "available" and "9.9.9" in card.title.text()
        assert "Rotate the picture." in card.notes.text() and "More detail" not in card.notes.text()
        assert card.btn_first.text() == "Update" and card.btn_second.text() == "Later"
        ctl.download(sync=True)
        assert card.state == "ready" and card.btn_first.text() == "Restart now"
        card.btn_second.click()  # "When I close TraceCMR"
        assert up.scheduled() == {"version": "9.9.9", "wheel": str(wheel), "relaunch": False,
                                  "notes": new.notes}
        assert not card.isVisible()
        ctl.check(manual=True, sync=True)  # asking again: it's waiting for the close
        assert card.state == "scheduled" and "installs when you close" in card.text.text()
        up.cancel_scheduled()

        answer["check"] = None
        ctl.check(manual=True, sync=True)
        assert card.state == "current" and "latest version" in card.text.text()
        answer["check"] = up.UpdateError("Couldn't check for updates (no network). Check the internet connection.")
        ctl.check(manual=True, sync=True)
        assert card.state == "check_failed" and card.btn_first.text() == "Try again"
        card.hide()
        ctl.check(manual=False, sync=True)  # the quiet check at start says nothing when offline
        assert not card.isVisible()

        monkeypatch.setattr(up, "can_patch", lambda u: ["numpy>=99 (have 2.0)"])
        answer["check"] = new
        ctl.check(manual=True, sync=True)
        assert card.state == "installer" and card.btn_first.text() == "Download"
        opened = []
        monkeypatch.setattr(ctl, "_open", lambda url: opened.append(url))
        card.btn_first.click()
        assert opened == ["https://example.org/TraceCMR-macOS.pkg"]

        monkeypatch.setattr(up, "install_kind", lambda: "dev")
        ctl.check(manual=True, sync=True)
        assert card.state == "dev" and "git pull" in card.text.text()

        ctl.show_result({"ok": True, "from": "0.3.4", "to": "9.9.9", "notes": ["Faster."]})
        assert card.state == "updated" and "9.9.9" in card.text.text()
        ctl.show_result({"ok": False, "from": "0.3.4", "to": "9.9.9",
                         "message": "The update to 9.9.9 couldn't be installed, so TraceCMR is still 0.3.4."})
        assert card.state == "update_failed" and "still 0.3.4" in card.text.text()
        assert "update_log.txt" in card.text.text()
    finally:
        up.cancel_scheduled()
        viewer.close()


def test_restart_now_saves_and_closes(qtbot, tmp_path, monkeypatch):
    from tracecmr import updater as up

    monkeypatch.setattr(up, "updates_dir", lambda: tmp_path / "updates")
    monkeypatch.setattr(up, "log_path", lambda: tmp_path / "update_log.txt")
    up.cancel_scheduled()
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.load_session(_phantom_session())
        saved = []
        monkeypatch.setattr(panel, "save_contours", lambda quiet=False, path=None: saved.append(quiet))
        monkeypatch.setattr(panel.autosave, "_current", lambda: ("changed", None))
        closed = []
        win = viewer.window._qt_window
        monkeypatch.setattr(win, "close", lambda: closed.append(1) or False)  # pretend something kept it open
        ctl = panel.updates
        ctl.up, ctl.wheel = up.Update(version="9.9.9"), tmp_path / "w.whl"
        ctl.restart_now()
        assert saved == [True] and closed == [1]
        assert up.scheduled()["relaunch"] is False  # kept open: install at the next close instead
    finally:
        up.cancel_scheduled()
        viewer.close()


def test_settings_and_help_offer_updates(qtbot):
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        texts = [a.text() for a in panel.btn_help.menu().actions()] + \
                [getattr(a, "_row_text", "") for a in panel.btn_help.menu().actions()]
        rows = panel.btn_help.menu().findChildren(type(panel.lbl_auto))
        assert any("Check for updates" in (r.text() or "") for r in rows) or any("Check for updates" in t
                                                                               for t in texts)
        assert panel.chk_updates.text() == "Check for updates when TraceCMR starts"
    finally:
        viewer.close()


def test_steadier_auto_lv_setting(qtbot, monkeypatch):
    """Settings -> "Steadier Auto LV" (on by default) turns the 8-orientation average on and off for the model
    process, and is remembered."""
    from tracecmr import autocontour as ac
    from tracecmr.app import _widget as W

    monkeypatch.setenv("TRACECMR_TTA", "1")
    reset = []
    monkeypatch.setattr(ac, "reset_workers", lambda: reset.append(1))
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        was = panel.chk_tta.isChecked()
        panel.chk_tta.setChecked(False)
        assert os.environ["TRACECMR_TTA"] == "0" and not ac.tta_enabled() and reset
        assert ac.epi_cutoff() == ac.EPI_CUTOFF  # the one-pass cutoff goes with one pass
        assert str(W._settings().value("auto/tta")).lower() in ("false", "0")
        panel.chk_tta.setChecked(True)
        assert ac.tta_enabled() and ac.epi_cutoff() == ac.EPI_CUTOFF_TTA
        panel.chk_tta.setChecked(was)
    finally:
        viewer.close()


def test_paravision_study_opens_as_one_stack(tmp_path):
    """Open scan on a Bruker ParaVision study folder: the short-axis scans (one per slice) open as one stack,
    the long-axis and localizer scans are listed; picking one scan's folder opens its study; picking a folder of
    studies lists them as mice."""
    from test_bruker import ORIENT, write_scan

    study = tmp_path / "mri" / "20260730_Mouse_1_44"
    normal = np.array(ORIENT[6:])
    for k, number in enumerate([10, 8, 9]):
        write_scan(study, number, f"SAX_S-{k}.0", np.full((16, 16, 1, 6), 100 * (k + 1)), slope=1.0,
                   positions=[list(np.array([7.4, 15.4, 9.9]) - k * normal)])
    write_scan(study, 16, "LAX_S0.0 4chamber", np.ones((16, 16, 1, 6), int), orient=[1, 0, 0, 0, 1, 0, 0, 0, 1])
    viewer = napari.Viewer(show=False)
    try:
        panel = TraceCMRWidget(viewer)
        panel._testing = True
        panel.open_path(study / "9")  # one scan's folder: the whole study opens
        s = panel.session
        assert s.image.shape == (16, 16, 3, 6) and s.separate_slices
        assert [float(s.image[0, 0, z, 0]) for z in range(3)] == [300.0, 200.0, 100.0]  # by position
        assert s.head_end() == "last" and s.order[0] == 2  # the slice toward the head is the base, from the scanner
        assert any(x.kind == "Long axis" for x in panel.study)
        panel.open_study(tmp_path / "mri")  # a folder of studies
        assert panel.series_panel.mice.count() == 1
        assert "ParaVision scan sessions" in panel.series_panel.lbl.text()
    finally:
        viewer.close()
