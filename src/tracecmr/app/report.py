"""One-click export: a workbook per mouse and a study table with one row per mouse.

Per mouse (in its results folder):
  <mouse>_TraceCMR_results.xlsx   tabs: Summary, Per slice, Wall segments, Notes
  images/                         bullseyes and the contoured slice montages (written by the panel)

Study table (in the results folder, shared by all mice):
  study_summary.xlsx              one row per mouse, one column per measure; exporting a mouse again
                                  updates its row instead of adding a second one

Excel files need openpyxl; without it the same tables are written as CSV files.
"""

from __future__ import annotations

import csv
import datetime as _dt
from pathlib import Path

import numpy as np

from .. import __version__

AHA_NAMES = {
    "basal": ["anterior", "anteroseptal", "inferoseptal", "inferior", "inferolateral", "anterolateral"],
    "mid": ["anterior", "anteroseptal", "inferoseptal", "inferior", "inferolateral", "anterolateral"],
    "apical": ["anterior", "septal", "inferior", "lateral"],
}
LEVEL_OFFSET = {"basal": 0, "mid": 6, "apical": 12}  # AHA numbers: basal 1-6, mid 7-12, apical 13-16

SUMMARY_FIELDS = [  # (key in the results, label, unit)
    ("edv_ul", "LV end-diastolic volume (EDV)", "uL"),
    ("esv_ul", "LV end-systolic volume (ESV)", "uL"),
    ("sv_ul", "LV stroke volume (SV)", "uL"),
    ("ef_pct", "LV ejection fraction (EF)", "%"),
    ("mass_ed_mg", "LV mass at ED", "mg"),
    ("mass_es_mg", "LV mass at ES", "mg"),
    ("rvedv_ul", "RV end-diastolic volume", "uL"),
    ("rvesv_ul", "RV end-systolic volume", "uL"),
    ("rvsv_ul", "RV stroke volume", "uL"),
    ("rvef_pct", "RV ejection fraction", "%"),
]


def _model_reference(session) -> str:
    """The paper to cite when the contours came from a published model (travels with every exported result)."""
    if "UCL" in (session.contoured_by or ""):
        from ..autocontour import UCL_CITATION

        return UCL_CITATION
    return ""


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return v
    return None if not np.isfinite(f) else round(f, 2)


def segment_rows(wall) -> list[dict]:
    """The 16 segments with AHA numbers and names (names only when segments are placed from the septum)."""
    if wall is None:
        return []
    aha = wall.ed.reference.startswith("septum")
    out = []
    for r in wall.rows():
        lv, s = r["level"], int(r["segment"])
        name = AHA_NAMES[lv][s - 1] if aha else f"sector {s} (image axis)"
        out.append({"aha": LEVEL_OFFSET[lv] + s if aha else None, "level": lv, "segment": s, "name": name,
                    "edwt_mm": _num(r["edwt_mm"]), "eswt_mm": _num(r["eswt_mm"]),
                    "thickening_pct": _num(r["thickening_pct"])})
    return out


def build(session, fn, wall, rv, notes) -> dict:
    """Everything the export writes, as plain tables."""
    row = fn.as_row()
    values = {**{k: row.get(k) for k in row}, **(rv or {})}
    summary = [(label, _num(values[k]), unit) for k, label, unit in SUMMARY_FIELDS if k in values]
    segs = segment_rows(wall)
    def mean(key):
        v = [r[key] for r in segs if isinstance(r[key], (int, float))]
        return round(float(np.mean(v)), 2) if v else None
    settings = [
        ("Mouse", session.subject_id),
        ("Scan file", session.source),
        ("Exported", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("TraceCMR version", __version__),
        ("Contoured by", session.contoured_by or ""),
        ("Auto-contour model reference", _model_reference(session)),
        ("Papillary muscles counted as", session.papillary or "not set"),
        ("ED frame (default)", session.ed_frame),
        ("ES frame (default)", session.es_frame),
        ("Per-slice ED/ES", "yes" if getattr(session, "separate_slices", False) else "no"),
        ("Slice order (base to apex)", " ".join(map(str, session.order))),
        ("Slices left out at ED", " ".join(map(str, sorted(set(session.order) - session.included["ED"]))) or "none"),
        ("Slices left out at ES", " ".join(map(str, sorted(set(session.order) - session.included["ES"]))) or "none"),
        ("Wall segment reference", wall.ed.reference if wall else "no wall results"),
        ("Septum set by hand on slices", " ".join(map(str, sorted(session.septum_manual))) or "none"),
        ("Wall thickness method", "centerline (perpendicular to the wall)"),
    ]
    per_slice = [{k: _num(v) for k, v in r.items()} for r in session.per_slice()]
    all_notes = list(notes) + list(fn.flags) + (list(wall.flags) if wall else [])
    study = {"mouse": session.subject_id, "exported": settings[2][1], "contoured_by": session.contoured_by or "",
             "model_reference": _model_reference(session),
             "papillary": session.papillary or "", "ed_frame": session.ed_frame, "es_frame": session.es_frame}
    study.update({k: _num(values[k]) for k, _, _ in SUMMARY_FIELDS if k in values})
    study.update({"n_slices_ed": row.get("n_slices_ed"), "n_slices_es": row.get("n_slices_es"),
                  "mean_edwt_mm": mean("edwt_mm"), "mean_thickening_pct": mean("thickening_pct"),
                  "segment_reference": wall.ed.reference if wall else ""})
    for r in segs:
        tag = f"{r['aha']}_{r['level']}_{r['name']}" if r["aha"] else f"{r['level']}_{r['segment']}"
        study[f"edwt_mm_{tag}"] = r["edwt_mm"]
        study[f"thickening_pct_{tag}"] = r["thickening_pct"]
    study["notes"] = len(all_notes)
    return {"summary": summary, "settings": settings, "per_slice": per_slice, "segments": segs,
            "notes": all_notes, "study": study}


# ---------- writing ----------

def _openpyxl():
    try:
        import openpyxl

        return openpyxl
    except ImportError:
        return None


def _style_header(ws, row, n):
    from openpyxl.styles import Font, PatternFill

    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="33414F")


def _widths(ws, widths):
    from openpyxl.utils import get_column_letter

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_mouse(path: Path, rep: dict) -> Path:
    """The mouse's workbook (or, without openpyxl, a folder of CSV files next to it). Returns what was written."""
    path = Path(path)
    xl = _openpyxl()
    if xl is None:
        return _write_mouse_csv(path.with_suffix(""), rep)
    wb = xl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Measure", "Value", "Unit"])
    _style_header(ws, 1, 3)
    for r in rep["summary"]:
        ws.append(list(r))
    ws.append([])
    ws.append(["Setting", "Value"])
    _style_header(ws, ws.max_row, 2)
    for r in rep["settings"]:
        ws.append(list(r))
    _widths(ws, [34, 40, 8])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Per slice")
    cols = ["slice", "edv_ul", "esv_ul", "ef_pct", "mass_ed_mg", "included"]
    ws.append(["Slice", "ED volume (uL)", "ES volume (uL)", "Slice EF (%)", "ED mass (mg)", "Counted"])
    _style_header(ws, 1, len(cols))
    for r in rep["per_slice"]:
        ws.append([("yes" if r.get(c) else "no") if c == "included" else r.get(c) for c in cols])
    _widths(ws, [8, 15, 15, 13, 13, 9])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Wall segments")
    ws.append(["AHA #", "Level", "Segment", "Name", "ED wall (mm)", "ES wall (mm)", "Thickening (%)"])
    _style_header(ws, 1, 7)
    for r in rep["segments"]:
        ws.append([r["aha"], r["level"], r["segment"], r["name"], r["edwt_mm"], r["eswt_mm"], r["thickening_pct"]])
    _widths(ws, [7, 9, 9, 24, 13, 13, 15])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Notes")
    ws.append(["Notes and warnings"])
    _style_header(ws, 1, 1)
    for n in rep["notes"] or ["none"]:
        ws.append([n])
    _widths(ws, [120])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _write_mouse_csv(folder: Path, rep: dict) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    with open(folder / "summary.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["measure", "value", "unit"])
        w.writerows(rep["summary"])
        w.writerow([])
        w.writerows(rep["settings"])
    for name, rows in (("per_slice", rep["per_slice"]), ("wall_segments", rep["segments"])):
        with open(folder / f"{name}.csv", "w", newline="", encoding="utf-8-sig") as fh:
            if rows:
                w = csv.DictWriter(fh, fieldnames=list(rows[0]))
                w.writeheader()
                w.writerows(rows)
    (folder / "notes.txt").write_text("\n".join(rep["notes"]) or "none", encoding="utf-8")
    return folder


def update_study(path: Path, row: dict, key: str) -> Path:
    """Add or update this mouse's row (matched on `key`) in the study table. Columns are added as needed.
    Without openpyxl a .csv is used instead. Raises PermissionError if the file is open in Excel."""
    path = Path(path)
    row = {"key": key, **row}
    xl = _openpyxl()
    if xl is None:
        path = path.with_suffix(".csv")
        rows = []
        if path.exists():
            with open(path, newline="", encoding="utf-8-sig") as fh:
                rows = list(csv.DictReader(fh))
        rows = [r for r in rows if r.get("key") != key] + [row]
        cols = list(dict.fromkeys(c for r in rows for c in r))
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)
        return path
    if path.exists():
        wb = xl.load_workbook(path)
        ws = wb["Mice"] if "Mice" in wb.sheetnames else wb.active
        header = [c.value for c in ws[1]]
    else:
        wb = xl.Workbook()
        ws = wb.active
        ws.title = "Mice"
        header = []
    for c in row:
        if c not in header:
            header.append(c)
            ws.cell(row=1, column=len(header), value=c)
    _style_header(ws, 1, len(header))
    target = None
    kcol = header.index("key") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=kcol).value == key:
            target = r
            break
    if target is None:
        target = ws.max_row + 1 if ws.max_row > 1 or ws.cell(row=2, column=1).value else 2
    for c, v in row.items():
        ws.cell(row=target, column=header.index(c) + 1, value=v)
    ws.freeze_panes = "C2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
